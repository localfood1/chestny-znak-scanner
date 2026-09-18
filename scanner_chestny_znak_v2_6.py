import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
import threading
import os
import csv
import ctypes
import sys

import fitz
import cv2
import numpy as np
import zxingcpp
from openpyxl import Workbook


# ================================================================
# DPI ДЛЯ WINDOWS
# ================================================================

try:
    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass


class ChestnyZnakScanner:

    # ============================================================
    # ЦВЕТА
    # ============================================================

    BG = "#111318"
    PANEL = "#181B21"
    PANEL_2 = "#20242C"
    BORDER = "#2C313B"

    TEXT = "#F1F3F5"
    TEXT_SECONDARY = "#9BA3AF"

    ACCENT = "#4C8DFF"
    ACCENT_HOVER = "#6EA3FF"

    SUCCESS = "#35C98B"
    DANGER = "#E85D6A"

    SELECTED = "#263D63"

    # ============================================================
    # ИНИЦИАЛИЗАЦИЯ
    # ============================================================

    def __init__(self, root):

        self.root = root

        self.root.geometry(
            "1250x750"
        )

        self.root.minsize(
            1000,
            600
        )

        self.root.configure(
            bg=self.BG
        )

        def resource_path(filename):
            if getattr(sys, "frozen", False):
                base_path = sys._MEIPASS
            else:
                base_path = os.path.dirname(os.path.abspath(__file__))

            return os.path.join(base_path, filename)

        root.iconbitmap(resource_path("icon_desktop.ico"))
        # --------------------------------------------------------
        # Данные
        # --------------------------------------------------------

        self.pdf_files = []
        self.results = []
        self.file_stats = {}
        self.seen_codes = set()
        self.scan_errors = []

        self.scanning = False
        self.stop_requested = False
        self.language = "ru"
        self.theme = "dark"
        self.settings_window = None
        self._detail_windows = []
        self._setup_localization()
        self.root.title(self.t("title"))

        # --------------------------------------------------------
        # Интерфейс
        # --------------------------------------------------------

        self.setup_styles()
        self.build_interface()

    # ============================================================
    # ЛОКАЛИЗАЦИЯ И ТЕМЫ
    # ============================================================

    def _setup_localization(self):
        self.languages = {
            "ru": "Русский", "en": "English", "kk": "Қазақша", "be": "Беларуская",
            "hy": "Հայերեն", "uz": "O‘zbekcha", "ky": "Кыргызча", "tg": "Тоҷикӣ",
            "ka": "ქართული", "tr": "Türkçe", "ce": "Нохчийн", "os": "Ирон"
        }
        # Значения хранятся по ключам, поэтому смена языка не затрагивает логику сканирования.
        ru = {
            "title":"Сканер маркировки — Честный знак","app":'Честный знак',"subtitle":'Сканер кодов маркировки из PDF',"ready":'Готов',
            "select":'📄  Выбрать PDF',"excel":'📊  Excel',"csv":'📋  CSV',"zpl":'🏷  Печать ZPL',"copy":'📋  Копировать',"details":'🔍  Подробности',"clear":'🗑  Очистить',"stop":'⛔  Остановить',
            "pdf_stat":'PDF-ФАЙЛОВ',"code_stat":'УНИКАЛЬНЫХ КОДОВ',"status_stat":'СТАТУС',"pdfs":"PDF-ФАЙЛЫ","file":'Файл',"codes":'Кодов',"marking":"КОДЫ МАРКИРОВКИ","number":'№',"page":'Страница',"full":'Полный код маркировки',
            "scanning":'Сканирование...',"stopping":'Остановка...',"stopped":"Остановлено","done":"Готово","select_files":"Выберите PDF-файлы","details_title":'Подробности кода',"details_subtitle":'Информация, распознанная из Data Matrix',"copy_gtin":'Копировать GTIN',"copy_serial":'Копировать серийный №',"copy_crypto":'Копировать криптохвост',"close":'Закрыть',
            "print_title":"Печать этикеток","print_question":"Напечатать выбранные коды ({count})?\nНажмите «Нет», чтобы напечатать все найденные коды.","label_size":"Размер этикетки","label_width":"Ширина этикетки, мм:","label_height":"Высота этикетки, мм:","confirm_print":"Подтвердите печать","send_print":"Отправить {count} этикеток на принтер:\n{printer}?\n\nПринтер должен поддерживать ZPL (например, Zebra).","sent":"Отправлено на печать: {count}","save_zpl":"Сохранить ZPL для термопринтера","zpl_file":"ZPL-файл","all_files":"Все файлы","zpl_saved":"ZPL сохранён","zpl_fallback":"Модуль pywin32 не установлен, поэтому этикетки сохранены в ZPL-файл. Его можно отправить на принтер Zebra.",
            "print_error":"Ошибка печати","save_error":"Ошибка сохранения ZPL","save_excel":"Сохранить результаты в Excel","excel_file":"Excel-файл","excel_saved":"Excel сохранён:\n{file}","excel_error":"Ошибка сохранения Excel","save_csv":"Сохранить результаты в CSV","csv_file":"CSV-файл","csv_saved":"CSV сохранён:\n{file}","csv_error":"Ошибка сохранения CSV","info":"Информация","error":"Ошибка","no_selection":"Выберите код.","copied":"Код скопирован","found":"Готово. Найдено кодов: {count}","selected":"Выбрано PDF: {count}","scan_error":"Ошибка сканирования",
            "gtin":"GTIN","serial":'Серийный номер',"crypto91":'Криптохвост AI 91',"crypto92":'Криптохвост AI 92',"crypto":'Криптохвост',"theme":"Тема","language":"Язык","light":"Светлая","dark":"Тёмная","settings":"Настройки","current_theme":"Текущая тема","close_settings":"Закрыть"
        }
        en = {
            "title":"Marking Scanner — Chestny Znak","app":"Chestny Znak","subtitle":"Marking code scanner from PDF","ready":"Ready","select":"📄  Select PDF","excel":'📊  Excel',"csv":'📋  CSV',"zpl":"🏷  Print ZPL","copy":"📋  Copy","details":"🔍  Details","clear":"🗑  Clear","stop":"⛔  Stop","pdf_stat":"PDF FILES","code_stat":"UNIQUE CODES","status_stat":"STATUS","pdfs":"PDF FILES","file":"File","codes":"Codes","marking":"MARKING CODES","number":"No.","page":"Page","full":"Full marking code","scanning":"Scanning...","stopping":"Stopping...","stopped":"Stopped","done":"Ready","select_files":"Select PDF files","details_title":"Code details","details_subtitle":"Information decoded from Data Matrix","copy_gtin":"Copy GTIN","copy_serial":"Copy serial number","copy_crypto":"Copy crypto tail","close":"Close","print_title":"Print labels","print_question":"Print selected codes ({count})?\nChoose No to print all found codes.","label_size":"Label size","label_width":"Label width, mm:","label_height":"Label height, mm:","confirm_print":"Confirm printing","send_print":"Send {count} labels to printer:\n{printer}?\n\nThe printer must support ZPL (for example, Zebra).","sent":"Sent to printer: {count}","save_zpl":"Save ZPL for thermal printer","zpl_file":"ZPL file","all_files":"All files","zpl_saved":"ZPL saved","zpl_fallback":"pywin32 is not installed, so the labels were saved to  ZPL file. You can send it to a Zebra printer.","print_error":"Print error","save_error":"ZPL save error","save_excel":"Save results to Excel","excel_file":"Excel file","excel_saved":"Excel saved:\n{file}","excel_error":"Excel save error","save_csv":"Save results to CSV","csv_file":"CSV file","csv_saved":"CSV saved:\n{file}","csv_error":"CSV save error","info":"Information","error":"Error","no_selection":"Select a code.","copied":"Code copied","found":"Done. Codes found: {count}","selected":"PDF selected: {count}","scan_error":"Scan error","gtin":"GTIN","serial":"Serial number","crypto91":"Crypto tail AI 91","crypto92":"Crypto tail AI 92","crypto":"Crypto tail","theme":"Theme","language":"Language","light":"Light","dark":"Dark","settings":"Settings","current_theme":"Current theme","close_settings":"Close"
        }
        self.translations={"ru":ru,"en":en}
        # Native UI translations for the requested languages.
        overrides={
          "kk":{"title":"Таңбалау сканері — Chestny Znak","app":"Chestny Znak","subtitle":"PDF файлдарынан таңбалау кодтарын сканерлеу","ready":"Дайын","select":"📄  PDF таңдау","copy":"📋  Көшіру","details":"🔍  Толығырақ","clear":"🗑  Тазалау","stop":"⛔  Тоқтату","pdf_stat":"PDF ФАЙЛДАРЫ","code_stat":"БІРЕГЕЙ КОДТАР","status_stat":"КҮЙІ","pdfs":"PDF ФАЙЛДАРЫ","file":'Файл',"codes":"Кодтар","marking":"ТАҢБАЛАУ КОДТАРЫ","page":"Бет","full":"Толық таңбалау коды","scanning":"Сканерлеу...","stopping":"Тоқтату...","stopped":"Тоқтатылды","done":"Дайын","language":"Тіл","theme":"Тақырып","light":"Ашық","dark":"Қараңғы","details_title":"Код мәліметтері","close":"Жабу","copy_gtin":"GTIN көшіру","copy_serial":"Сериялық нөмірді көшіру","copy_crypto":"Криптоқұйрықты көшіру"},
          "be":{"title":"Сканер маркіроўкі — Честны знак","app":"Честны знак","subtitle":"Сканер кодаў маркіроўкі з PDF","ready":"Гатова","select":"📄  Выбраць PDF","copy":"📋  Капіяваць","details":"🔍  Падрабязнасці","clear":"🗑  Ачысціць","stop":"⛔  Спыніць","pdf_stat":"PDF-ФАЙЛЫ","code_stat":"УНІКАЛЬНЫХ КОДАЎ","status_stat":'СТАТУС',"pdfs":"PDF-ФАЙЛЫ","file":'Файл',"codes":"Кодаў","marking":"КОДЫ МАРКІРОЎКІ","page":"Старонка","full":"Поўны код маркіроўкі","scanning":"Сканаванне...","stopping":"Спыненне...","stopped":"Спынена","done":"Гатова","language":"Мова","theme":"Тэма","light":"Светлая","dark":"Цёмная","details_title":"Падрабязнасці кода","close":"Закрыць","copy_gtin":"Капіяваць GTIN","copy_serial":"Капіяваць серыйны нумар","copy_crypto":"Капіяваць крыптахвост"},
          "hy":{"title":"Նշագրման սկաներ — Չեստնի Զնակ","app":"Չեստնի Զնակ","subtitle":"Նշագրման կոդերի սկանավորում PDF-ից","ready":"Պատրաստ է","select":"📄  Ընտրել PDF","copy":"📋  Պատճենել","details":"🔍  Մանրամասներ","clear":"🗑  Մաքրել","stop":"⛔  Կանգնեցնել","pdf_stat":"PDF ՖԱՅԼԵՐ","code_stat":"ԵԶԱԿԻ ԿՈԴԵՐ","status_stat":"ԿԱՐԳԱՎԻՃԱԿ","pdfs":"PDF ՖԱՅԼԵՐ","file":"Ֆայլ","codes":"Կոդեր","marking":"ՆՇԱԳՐՄԱՆ ԿՈԴԵՐ","page":"Էջ","full":"Ամբողջական նշագրման կոդ","scanning":"Սկանավորում...","stopping":"Կանգնեցում...","stopped":"Կանգնեցված է","done":"Պատրաստ է","language":"Լեզու","theme":"Թեմա","light":"Բաց","dark":"Մուգ","details_title":"Կոդի մանրամասներ","close":"Փակել","copy_gtin":"Պատճենել GTIN","copy_serial":"Պատճենել սերիական համարը","copy_crypto":"Պատճենել կրիպտոպոչը"},
          "tr":{"title":"Etiket Tarayıcı — Chestny Znak","app":"Chestny Znak","subtitle":"PDF'den işaretleme kodu tarayıcı","ready":"Hazır","select":"📄  PDF Seç","copy":"📋  Kopyala","details":"🔍  Ayrıntılar","clear":"🗑  Temizle","stop":"⛔  Durdur","pdf_stat":"PDF DOSYALARI","code_stat":"BENZERSİZ KODLAR","status_stat":"DURUM","pdfs":"PDF DOSYALARI","file":"Dosya","codes":"Kodlar","marking":"İŞARETLEME KODLARI","page":"Sayfa","full":"Tam işaretleme kodu","scanning":"Taranıyor...","stopping":"Durduruluyor...","stopped":"Durduruldu","done":"Hazır","language":"Dil","theme":"Tema","light":"Açık","dark":"Koyu","details_title":"Kod ayrıntıları","close":"Kapat","copy_gtin":"GTIN'i kopyala","copy_serial":"Seri numarasını kopyala","copy_crypto":"Kripto kuyruğunu kopyala"},
          "uz":{"title":"Markirovka skaneri — Chestny Znak","app":"Chestny Znak","subtitle":"PDF dan markirovka kodlarini skanerlash","ready":"Tayyor","select":"📄  PDF tanlash","copy":"📋  Nusxalash","details":"🔍  Tafsilotlar","clear":"🗑  Tozalash","stop":"⛔  To‘xtatish","pdf_stat":"PDF FAYLLAR","code_stat":"NOYOB KODLAR","status_stat":"HOLAT","pdfs":"PDF FAYLLAR","file":"Fayl","codes":"Kodlar","marking":"MARKIROVKA KODLARI","page":"Sahifa","full":"To‘liq markirovka kodi","scanning":"Skanerlanmoqda...","stopping":"To‘xtatilmoqda...","stopped":"To‘xtatildi","done":"Tayyor","language":"Til","theme":"Mavzu","light":"Yorug‘","dark":"Qorong‘i","details_title":"Kod tafsilotlari","close":"Yopish","copy_gtin":"GTIN nusxalash","copy_serial":"Seriya raqamini nusxalash","copy_crypto":"Kripto qismini nusxalash"},
          "ky":{"title":"Маркировка сканери — Chestny Znak","app":"Chestny Znak","subtitle":"PDF файлдарынан маркировка коддорун сканерлөө","ready":"Даяр","select":"📄  PDF тандоо","copy":"📋  Көчүрүү","details":"🔍  Толук маалымат","clear":"🗑  Тазалоо","stop":"⛔  Токтотуу","pdf_stat":"PDF ФАЙЛДАРЫ","code_stat":"БИРЕГЕЙ КОДДОР","status_stat":"АБАЛЫ","pdfs":"PDF ФАЙЛДАРЫ","file":'Файл',"codes":"Коддор","marking":"МАРКИРОВКА КОДДОРУ","page":"Барак","full":"Толук маркировка коду","scanning":"Сканерлөө...","stopping":"Токтотулууда...","stopped":"Токтотулду","done":"Даяр","language":"Тил","theme":"Тема","light":"Ачык","dark":"Караңгы","details_title":"Код тууралуу маалымат","close":"Жабуу","copy_gtin":"GTIN көчүрүү","copy_serial":"Сериялык номерди көчүрүү","copy_crypto":"Крипто бөлүгүн көчүрүү"},
          "tg":{"title":"Сканери тамғагузорӣ — Chestny Znak","app":"Chestny Znak","subtitle":"Сканкунии рамзҳои тамғагузорӣ аз PDF","ready":"Тайёр","select":"📄  Интихоби PDF","copy":"📋  Нусхабардорӣ","details":"🔍  Тафсилот","clear":"🗑  Пок кардан","stop":"⛔  Қатъ кардан","pdf_stat":"ФАЙЛҲОИ PDF","code_stat":"РАМЗҲОИ НОДИР","status_stat":"ҲОЛАТ","pdfs":"ФАЙЛҲОИ PDF","file":'Файл',"codes":"Рамзҳо","marking":"РАМЗҲОИ ТАМҒАГУЗОРӢ","page":"Саҳифа","full":"Рамзи пурраи тамғагузорӣ","scanning":"Сканкунӣ...","stopping":"Қатъкунӣ...","stopped":"Қатъ шуд","done":"Тайёр","language":"Забон","theme":"Мавзӯъ","light":"Равшан","dark":"Торик","details_title":"Тафсилоти рамз","close":"Пӯшидан","copy_gtin":"Нусхабардории GTIN","copy_serial":"Нусхабардории рақами силсилавӣ","copy_crypto":"Нусхабардории қисми крипто"},
          "ka":{"title":"მარკირების სკანერი — Chestny Znak","app":"Chestny Znak","subtitle":"მარკირების კოდების სკანერი PDF-დან","ready":"მზადაა","select":"📄  PDF-ის არჩევა","copy":"📋  კოპირება","details":"🔍  დეტალები","clear":"🗑  გასუფთავება","stop":"⛔  გაჩერება","pdf_stat":"PDF ფაილები","code_stat":"უნიკალური კოდები","status_stat":"სტატუსი","pdfs":"PDF ფაილები","file":"ფაილი","codes":"კოდები","marking":"მარკირების კოდები","page":"გვერდი","full":"მარკირების სრული კოდი","scanning":"სკანირება...","stopping":"შეჩერება...","stopped":"შეჩერდა","done":"მზადაა","language":"ენა","theme":"თემა","light":"ღია","dark":"მუქი","details_title":"კოდის დეტალები","close":"დახურვა","copy_gtin":"GTIN-ის კოპირება","copy_serial":"სერიული ნომრის კოპირება","copy_crypto":"კრიპტო ნაწილის კოპირება"},
          "ce":{"title":"Маркировкин сканер — Chestny Znak","app":"Chestny Znak","subtitle":"PDF тӀера маркировкин кодаш хьовса","ready":"Кхиам болу","select":"📄  PDF талла","copy":"📋  ДӀаяздар","details":"🔍  Бакъо","clear":"🗑  Чистаяккха","stop":"⛔  ТоӀа дар","pdf_stat":"PDF-ФАЙЛАШ","code_stat":"ЙИШ Ю КОДАШ","status_stat":'СТАТУС',"pdfs":"PDF-ФАЙЛАШ","file":'Файл',"codes":"Кодаш","marking":"МАРКИРОВКИН КОДАШ","page":"Оьрта","full":"Маркировкин юккъера код","scanning":"Хьовсуш ду...","stopping":"ТоӀа дар...","stopped":"ТоӀа дина","done":"Кхиам болу","language":"Мотт","theme":"Тема","light":"Жуьра","dark":"Хьалха","details_title":"Кодан бакъо","close":"ЧӀагӀо","copy_gtin":"GTIN дӀаяздар","copy_serial":"Серийн номер дӀаяздар","copy_crypto":"Криптохвост дӀаяздар"},
          "os":{"title": "Маркировкӕйы сканер — Chestny Znak","app": "Chestny Znak","subtitle": "Маркировкӕйы кодтӕ PDF-файлты скан кӕнын","ready": "Дӕрӕн","select": "📄  PDF рауадзын","copy": "📋  Равдисын","details": "🔍  Фӕлтӕр","clear": "🗑  Раст кӕнын","stop": "⛔  Стоп","pdf_stat": "PDF-ФАЙЛТӔ","code_stat": "ЦЫБЫР КОДТӔ","status_stat": "СТАТУС","pdfs": "PDF-ФАЙЛТӔ","file": "Файл","codes": "Кодтӕ","marking": "МАРКИРОВКӔЙЫ КОДТӔ","page": "Фӕрс","full": "Маркировкӕйы ӕмбӕл код","scanning": "Скан кӕны...","stopping": "Стоп кӕны...","stopped": "Стоп кӕдтӕн","done": "Дӕрӕн","select_files": "Равдис PDF-файлты","details_title": "Кодӕйы фӕлтӕр","details_subtitle": "Data Matrix-ы равдистӕйы информация","copy_gtin": "GTIN равдисын","copy_serial": "Серийн номер равдисын","copy_crypto": "Криптохвост равдисын","close": "Закрыть","print_title": "Этикеткӕтӕйы фӕлӕхст","print_question": "Равдис выбран кодтӕ ({count})?\n«Нӕ» бас, цӕмӕн ӕй ӕрцӕу кодтӕйы фӕлӕхст.","label_size": "Этикеткӕйы ӕндӕр","label_width": "Этикеткӕйы ӕргъӕ, мм:","label_height": "Этикеткӕйы ӕмбӕрц, мм:","confirm_print": "Фӕлӕхстӕйы тӕрхон","send_print": "{count} этикеткӕ принтермӕ равдисын:\n{printer}?\n\nПринтер уӕлдай ZPL хъӕуы.","sent": "Фӕлӕхстмӕ равдис: {count}","save_zpl": "ZPL термопринтеры фӕстӕгмӕ ӕрвитын","zpl_file": "ZPL-файл","all_files": "Ӕппӕт файлтӕ","zpl_saved": "ZPL сақ кӕдтӕн","zpl_fallback": "pywin32 нӕ у, уый тыххӕй этикеткӕтӕ ZPL-файлмӕ сақ кӕдтӕн.","print_error": "Фӕлӕхсты хӕс","save_error": "ZPL-йы сақ кӕныны хӕс","save_excel": "Нӕтижӕнтӕ Excel-мӕ сақ кӕнын","excel_file": "Excel-файл","excel_saved": "Excel сақ кӕдтӕн:\n{file}","excel_error": "Excel-йы хӕс","save_csv": "Нӕтижӕнтӕ CSV-мӕ сақ кӕнын","csv_file": "CSV-файл","csv_saved": "CSV сақ кӕдтӕн:\n{file}","csv_error": "CSV-йы хӕс","info": "Информаци","error": "Хӕс","no_selection": "Равдис код.","copied": "Код равдисӕй ӕрцӕу","found": "Дӕрӕн. Нӕмӕ кодтӕ: {count}","selected": "PDF равдистӕ: {count}","scan_error": "Скан кӕныны хӕс","gtin": "GTIN","serial": "Серийн номер","crypto91": "Криптохвост AI 91","crypto92": "Криптохвост AI 92","crypto": "Криптохвост","theme": "Темӕ","language": "Ӕвзаг","light": "Рауаг","dark": "Тъӕпп","settings": "Настройкӕ","current_theme": "Ныртӕккӕйы темӕ","close_settings": "Закрыть"
        }

        }

        for lang, over in overrides.items():
            d=dict(en); d.update(over); self.translations[lang]=d

        # Базовые названия настроек для всех дополнительных языков.
        # Основные элементы интерфейса выше уже переведены отдельно.
        settings_names = {
            "kk": {"settings":"Баптаулар","current_theme":"Ағымдағы тақырып","close_settings":"Жабу"},
            "be": {"settings":"Налады","current_theme":"Бягучая тэма","close_settings":"Закрыць"},
            "hy": {"settings":"Կարգավորումներ","current_theme":"Ընթացիկ թեմա","close_settings":"Փակել"},
            "uz": {"settings":"Sozlamalar","current_theme":"Joriy mavzu","close_settings":"Yopish"},
            "ky": {"settings":"Жөндөөлөр","current_theme":"Учурдагы тема","close_settings":"Жабуу"},
            "tg": {"settings":"Танзимот","current_theme":"Мавзӯи ҷорӣ","close_settings":"Пӯшидан"},
            "ka": {"settings":"პარამეტრები","current_theme":"მიმდინარე თემა","close_settings":"დახურვა"},
            "tr": {"settings":"Ayarlar","current_theme":"Mevcut tema","close_settings":"Kapat"},
            "ce": {"settings":"Хьайоьрмаш","current_theme":"Хьалха йоцу тема","close_settings":"ЧӀагӀо"},
            "os": {"settings": "Настройкӕ","current_theme": "Ныртӕккӕйы темӕ","close_settings": "Закрыть"
         },
        }
        for lang, values in settings_names.items():
            self.translations[lang].update(values)

        # Дополнительные элементы интерфейса, которые должны переводиться сразу.
        core = {
            "kk": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  ZPL басып шығару", "number":"№", "serial_ai21":"Сериялық нөмір (AI 21)", "copied_full":"Толық код көшірілді", "copied_gtin":"GTIN көшірілді", "copied_serial":"Сериялық нөмір көшірілді", "copied_crypto":"Криптоқұйрық көшірілді"},
            "be": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  Друк ZPL", "number":"№", "serial_ai21":"Серыйны нумар (AI 21)", "copied_full":"Поўны код скапіяваны", "copied_gtin":"GTIN скапіяваны", "copied_serial":"Серыйны нумар скапіяваны", "copied_crypto":"Крыптахвост скапіяваны"},
            "hy": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  Տպել ZPL", "number":"№", "serial_ai21":"Սերիական համար (AI 21)", "copied_full":"Ամբողջական կոդը պատճենված է", "copied_gtin":"GTIN-ը պատճենված է", "copied_serial":"Սերիական համարը պատճենված է", "copied_crypto":"Կրիպտոպոչը պատճենված է"},
            "uz": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  ZPL chop etish", "number":"№", "serial_ai21":"Seriya raqami (AI 21)", "copied_full":"To‘liq kod nusxalandi", "copied_gtin":"GTIN nusxalandi", "copied_serial":"Seriya raqami nusxalandi", "copied_crypto":"Kripto qismi nusxalandi"},
            "ky": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  ZPL басып чыгаруу", "number":"№", "serial_ai21":"Сериялык номер (AI 21)", "copied_full":"Толук код көчүрүлдү", "copied_gtin":"GTIN көчүрүлдү", "copied_serial":"Сериялык номер көчүрүлдү", "copied_crypto":"Крипто бөлүгү көчүрүлдү"},
            "tg": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  Чопи ZPL", "number":"№", "serial_ai21":"Рақами силсилавӣ (AI 21)", "copied_full":"Рамзи пурра нусхабардорӣ шуд", "copied_gtin":"GTIN нусхабардорӣ шуд", "copied_serial":"Рақами силсилавӣ нусхабардорӣ шуд", "copied_crypto":"Қисми крипто нусхабардорӣ шуд"},
            "ka": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  ZPL ბეჭდვა", "number":"№", "serial_ai21":"სერიული ნომერი (AI 21)", "copied_full":"სრული კოდი დაკოპირდა", "copied_gtin":"GTIN დაკოპირდა", "copied_serial":"სერიული ნომერი დაკოპირდა", "copied_crypto":"კრიპტო ნაწილი დაკოპირდა"},
            "tr": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  ZPL Yazdır", "number":"No.", "serial_ai21":"Seri numarası (AI 21)", "copied_full":"Tam kod kopyalandı", "copied_gtin":"GTIN kopyalandı", "copied_serial":"Seri numarası kopyalandı", "copied_crypto":"Kripto kısmı kopyalandı"},
            "ce": {"excel":"📊  Excel", "csv":"📋  CSV", "zpl":"🏷  ZPL чӀагӀорг", "number":"№", "serial_ai21":"Серийн номер (AI 21)", "copied_full":"Маркировкин юккъера код дӀаяздина", "copied_gtin":"GTIN дӀаяздина", "copied_serial":"Серийн номер дӀаяздина", "copied_crypto":"Криптохвост дӀаяздина"},
            "os": {"excel": "📊  Excel","csv": "📋  CSV","zpl": "🏷  ZPL фӕлӕхст","number": "№","serial_ai21": "Серийн номер (AI 21)","copied_full": "Ӕмбӕл код равдисӕй ӕрцӕу","copied_gtin": "GTIN равдисӕй ӕрцӕу","copied_serial": "Серийн номер равдисӕй ӕрцӕу","copied_crypto": "Криптохвост равдисӕй ӕрцӕу"
        },

        }
        for lang, values in core.items():
            self.translations[lang].update(values)
        self.translations["ru"].update({
            "serial_ai21":"Серийный номер (AI 21)",
            "copied_full":"Полный код скопирован",
            "copied_gtin":"GTIN скопирован",
            "copied_serial":"Серийный номер скопирован",
            "copied_crypto":"Криптохвост скопирован",
        })
        self.translations["en"].update({
            "serial_ai21":"Serial number (AI 21)",
            "copied_full":"Full code copied",
            "copied_gtin":"GTIN copied",
            "copied_serial":"Serial number copied",
            "copied_crypto":"Crypto tail copied",
        })

    def t(self,key,**kwargs):
        text=self.translations.get(self.language,self.translations["ru"]).get(key,self.translations["en"].get(key,key))
        return text.format(**kwargs) if kwargs else text

    def _change_language(self, lang):
        if lang not in self.languages:
            return
        self.language = lang
        self._refresh_localized_ui()
        self._refresh_settings_window()
        self._refresh_detail_windows()

    def _toggle_theme(self):
        self.theme = "light" if self.theme == "dark" else "dark"
        self._apply_theme()

    def _center_window(self, window, width, height):
        window.update_idletasks()
        screen_w = window.winfo_screenwidth()
        screen_h = window.winfo_screenheight()
        x = max(0, (screen_w - width) // 2)
        y = max(0, (screen_h - height) // 2)
        window.geometry(f"{width}x{height}+{x}+{y}")

    def _open_settings(self):
        if getattr(self, "settings_window", None) is not None:
            try:
                if self.settings_window.winfo_exists():
                    self.settings_window.lift()
                    self.settings_window.focus_force()
                    return
            except Exception:
                pass

        win = tk.Toplevel(self.root)
        self.settings_window = win
        win.title(self.t("settings"))
        win.geometry("380x250")
        win.minsize(340, 220)
        win.configure(bg=self.PANEL)
        win.transient(self.root)

        frame = tk.Frame(win, bg=self.PANEL)
        frame.pack(fill="both", expand=True, padx=22, pady=20)
        self.settings_frame = frame

        self.settings_language_label = tk.Label(
            frame,
            text=self.t("language") + ":",
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 10, "bold")
        )
        self.settings_language_label.pack(anchor="w", pady=(0, 6))

        self.settings_language_var = tk.StringVar(value=self.languages[self.language])
        self.settings_language_combo = ttk.Combobox(
            frame,
            textvariable=self.settings_language_var,
            values=list(self.languages.values()),
            state="readonly",
            width=28
        )
        self.settings_language_combo.pack(fill="x", pady=(0, 18))
        self.settings_language_combo.bind(
            "<<ComboboxSelected>>",
            lambda e: self._change_language(
                next(k for k, v in self.languages.items()
                     if v == self.settings_language_var.get())
            )
        )

        self.settings_theme_label = tk.Label(
            frame,
            text=self.t("current_theme") + ":",
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 10, "bold")
        )
        self.settings_theme_label.pack(anchor="w", pady=(0, 6))

        self.settings_theme_button = ttk.Button(
            frame,
            command=self._toggle_theme,
            style="Dark.TButton"
        )
        self.settings_theme_button.pack(fill="x", pady=(0, 18))

        self.settings_close_button = ttk.Button(
            frame,
            text=self.t("close_settings"),
            command=self._close_settings,
            style="Dark.TButton"
        )
        self.settings_close_button.pack(fill="x")

        win.protocol("WM_DELETE_WINDOW", self._close_settings)
        self._refresh_settings_window()
        self._center_window(win, 380, 250)
        win.lift()
        win.focus_force()

    def _close_settings(self):
        try:
            if self.settings_window is not None and self.settings_window.winfo_exists():
                self.settings_window.destroy()
        except Exception:
            pass
        self.settings_window = None

    def _refresh_settings_window(self):
        win = getattr(self, "settings_window", None)
        try:
            if win is None or not win.winfo_exists():
                return
            win.title(self.t("settings"))
            win.configure(bg=self.PANEL)
            self.settings_frame.configure(bg=self.PANEL)
            self.settings_language_label.configure(
                text=self.t("language") + ":", bg=self.PANEL, fg=self.TEXT
            )
            self.settings_language_var.set(self.languages[self.language])
            self.settings_theme_label.configure(
                text=self.t("current_theme") + ":", bg=self.PANEL, fg=self.TEXT
            )
            self.settings_theme_button.configure(text=self._theme_button_text())
            self.settings_close_button.configure(text=self.t("close_settings"))
        except Exception:
            pass

    def _theme_button_text(self):
        return "🌙  " + self.t("dark") if self.theme == "dark" else "☀  " + self.t("light")

    def _apply_theme(self):
        if self.theme == "dark":
            self.BG="#111318"; self.PANEL="#181B21"; self.PANEL_2="#20242C"; self.BORDER="#2C313B"; self.TEXT="#F1F3F5"; self.TEXT_SECONDARY="#9BA3AF"; self.SELECTED="#263D63"
        else:
            self.BG="#F4F6F8"; self.PANEL="#FFFFFF"; self.PANEL_2="#E8ECF1"; self.BORDER="#D0D7DE"; self.TEXT="#1F2328"; self.TEXT_SECONDARY="#5F6B7A"; self.SELECTED="#D9E7FF"

        self.setup_styles()
        self.root.configure(bg=self.BG)

        for w in getattr(self, "_bg_frames", []):
            try:
                w.configure(bg=self.BG)
            except Exception:
                pass
        for w in getattr(self, "_panel_frames", []):
            try:
                w.configure(bg=self.PANEL)
            except Exception:
                pass
        for w in getattr(self, "_panel_labels", []):
            try:
                w.configure(
                    bg=self.PANEL,
                    fg=self.TEXT if getattr(w, "_is_value", False) else self.TEXT_SECONDARY
                )
            except Exception:
                pass

        if hasattr(self, "header_status"):
            self.header_status.configure(bg=self.BG, fg=self.SUCCESS)
        if hasattr(self, "toolbar_inner"):
            self.toolbar_inner.configure(bg=self.PANEL)

        self._refresh_localized_ui()
        self._refresh_settings_window()
        self._refresh_detail_windows()

    def _refresh_localized_ui(self):
        self.root.title(self.t("title"))
        self.app_title_label.configure(text=self.t("app"))
        self.subtitle_label.configure(text=self.t("subtitle"))
        self.select_button.configure(text=self.t("select"))
        self.excel_button.configure(text=self.t("excel"))
        self.csv_button.configure(text=self.t("csv"))
        self.print_button.configure(text=self.t("zpl"))
        self.copy_button.configure(text=self.t("copy"))
        self.details_button.configure(text=self.t("details"))
        self.clear_button.configure(text=self.t("clear"))
        self.stop_button.configure(text=self.t("stop"))
        self.pdf_card_title.configure(text=self.t("pdf_stat"))
        self.code_card_title.configure(text=self.t("code_stat"))
        self.status_card_title.configure(text=self.t("status_stat"))
        self.file_frame.configure(text=f"  {self.t('pdfs')}  ")
        self.file_tree.heading("#0", text=self.t("file"))
        self.file_tree.heading("count", text=self.t("codes"))
        self.code_frame.configure(text=f"  {self.t('marking')}  ")
        self.code_tree.heading("number", text=self.t("number"))
        self.code_tree.heading("page", text=self.t("page"))
        self.code_tree.heading("code", text=self.t("full"))
        if self.scanning:
            current_status = self.t("scanning")
            header_text = "● " + self.t("scanning")
        elif self.stop_requested:
            current_status = self.t("stopped")
            header_text = "● " + self.t("stopped")
        elif self.results:
            current_status = self.t("found", count=len(self.results))
            header_text = "● " + self.t("done")
        else:
            current_status = self.t("ready")
            header_text = "● " + self.t("ready")
        self.header_status.configure(text=header_text)
        self.status_label.configure(text=current_status)
        self.status_value.configure(text=current_status[:22])
        self.settings_button.configure(text="⚙  " + self.t("settings"))
        self._refresh_settings_window()

    def _refresh_detail_windows(self):
        alive = []
        for refs in getattr(self, "_detail_windows", []):
            try:
                if not refs["window"].winfo_exists():
                    continue
                win = refs["window"]
                win.title(self.t("details_title"))
                win.configure(bg=self.BG)
                refs["header"].configure(bg=self.BG)
                refs["title"].configure(text=self.t("details_title"), bg=self.BG, fg=self.TEXT)
                refs["subtitle"].configure(text=self.t("details_subtitle"), bg=self.BG, fg=self.TEXT_SECONDARY)
                refs["frame"].configure(bg=self.PANEL, highlightbackground=self.BORDER)
                refs["content"].configure(bg=self.PANEL)
                refs["full_label"].configure(text=self.t("full"), bg=self.PANEL, fg=self.TEXT)
                refs["text_frame"].configure(bg=self.PANEL_2, highlightbackground=self.BORDER)
                refs["full_text"].configure(
                    bg=self.PANEL_2, fg=self.TEXT, insertbackground=self.TEXT,
                    selectbackground=self.SELECTED
                )
                refs["buttons_frame"].configure(bg=self.PANEL)
                refs["copy_gtin"].configure(text=self.t("copy_gtin"))
                refs["copy_serial"].configure(text=self.t("copy_serial"))
                refs["copy_crypto"].configure(text=self.t("copy_crypto"))
                refs["close"].configure(text=self.t("close"))
                for label, entry, key in refs["fields"]:
                    label.configure(text=self.t(key), bg=self.PANEL, fg=self.TEXT)
                    entry.configure(
                        bg=self.PANEL_2, fg=self.TEXT,
                        insertbackground=self.TEXT,
                        readonlybackground=self.PANEL_2,
                        highlightbackground=self.BORDER
                    )
                alive.append(refs)
            except Exception:
                pass
        self._detail_windows = alive

    # ============================================================
    # СТИЛИ
    # ============================================================

    def setup_styles(self):

        style = ttk.Style()

        try:
            style.theme_use("clam")
        except Exception:
            pass

        # --------------------------------------------------------
        # Основные элементы
        # --------------------------------------------------------

        style.configure(
            ".",
            background=self.BG,
            foreground=self.TEXT,
            font=("Segoe UI", 10)
        )

        style.configure(
            "TFrame",
            background=self.BG
        )

        style.configure(
            "TLabel",
            background=self.BG,
            foreground=self.TEXT
        )

        # --------------------------------------------------------
        # Заголовок
        # --------------------------------------------------------

        style.configure(
            "Title.TLabel",
            background=self.BG,
            foreground=self.TEXT,
            font=("Segoe UI", 18, "bold")
        )

        style.configure(
            "Subtitle.TLabel",
            background=self.BG,
            foreground=self.TEXT_SECONDARY,
            font=("Segoe UI", 9)
        )

        # --------------------------------------------------------
        # LabelFrame
        # --------------------------------------------------------

        style.configure(
            "Dark.TLabelframe",
            background=self.PANEL,
            foreground=self.TEXT,
            bordercolor=self.BORDER,
            relief="solid",
            borderwidth=1
        )

        style.configure(
            "Dark.TLabelframe.Label",
            background=self.PANEL,
            foreground=self.TEXT,
            font=("Segoe UI", 10, "bold")
        )

        # --------------------------------------------------------
        # Обычные кнопки
        # --------------------------------------------------------

        style.configure(
            "Dark.TButton",
            background=self.PANEL_2,
            foreground=self.TEXT,
            bordercolor=self.BORDER,
            borderwidth=1,
            padding=(13, 8),
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Dark.TButton",
            background=[
                ("active", "#2A303A"),
                ("pressed", "#303743"),
                ("disabled", "#15171B")
            ],
            foreground=[
                ("disabled", "#555B66")
            ]
        )

        # --------------------------------------------------------
        # Главная кнопка
        # --------------------------------------------------------

        style.configure(
            "Accent.TButton",
            background=self.ACCENT,
            foreground="white",
            borderwidth=0,
            padding=(15, 9),
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Accent.TButton",
            background=[
                ("active", self.ACCENT_HOVER),
                ("pressed", self.ACCENT)
            ]
        )

        # --------------------------------------------------------
        # Кнопка остановки
        # --------------------------------------------------------

        style.configure(
            "Danger.TButton",
            background=self.DANGER,
            foreground="white",
            borderwidth=0,
            padding=(13, 8),
            font=("Segoe UI", 9, "bold")
        )

        # --------------------------------------------------------
        # Treeview
        # --------------------------------------------------------

        style.configure(
            "Dark.Treeview",
            background=self.PANEL,
            fieldbackground=self.PANEL,
            foreground=self.TEXT,
            bordercolor=self.BORDER,
            rowheight=34,
            font=("Segoe UI", 9)
        )

        style.map(
            "Dark.Treeview",
            background=[
                ("selected", self.SELECTED)
            ],
            foreground=[
                ("selected", "white")
            ]
        )

        style.configure(
            "Dark.Treeview.Heading",
            background=self.PANEL_2,
            foreground=self.TEXT_SECONDARY,
            bordercolor=self.BORDER,
            relief="flat",
            padding=(8, 9),
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Dark.Treeview.Heading",
            background=[
                ("active", "#2A303A")
            ]
        )

        # --------------------------------------------------------
        # Scrollbars
        # --------------------------------------------------------

        style.configure(
            "Dark.Vertical.TScrollbar",
            background=self.PANEL_2,
            troughcolor=self.PANEL,
            bordercolor=self.PANEL,
            arrowcolor=self.TEXT_SECONDARY
        )

        style.configure(
            "Dark.Horizontal.TScrollbar",
            background=self.PANEL_2,
            troughcolor=self.PANEL,
            bordercolor=self.PANEL,
            arrowcolor=self.TEXT_SECONDARY
        )

        # --------------------------------------------------------
        # Progressbar
        # --------------------------------------------------------

        style.configure(
            "Dark.Horizontal.TProgressbar",
            troughcolor=self.PANEL_2,
            background=self.ACCENT,
            bordercolor=self.PANEL_2,
            lightcolor=self.ACCENT,
            darkcolor=self.ACCENT
        )

    # ============================================================
    # ИНТЕРФЕЙС
    # ============================================================

    def build_interface(self):

        # ========================================================
        # ШАПКА
        # ========================================================

        header = tk.Frame(
            self.root,
            bg=self.BG
        )

        header.pack(
            fill="x",
            padx=22,
            pady=(18, 8)
        )

        title_frame = tk.Frame(
            header,
            bg=self.BG
        )

        self._bg_frames = [header, title_frame]

        title_frame.pack(
            side="left"
        )

        self.app_title_label = ttk.Label(
            title_frame,
            text=self.t("app"),
            style="Title.TLabel"
        )
        self.app_title_label.pack(anchor="w")

        self.subtitle_label = ttk.Label(
            title_frame,
            text=self.t("subtitle"),
            style="Subtitle.TLabel"
        )
        self.subtitle_label.pack(anchor="w", pady=(2, 0))

        self.header_status = tk.Label(
            header,
            text="● "+self.t("ready"),
            bg=self.BG,
            fg=self.SUCCESS,
            font=("Segoe UI", 10, "bold")
        )

        self.header_status.pack(
            side="right",
            pady=5
        )

        # ========================================================
        # ПАНЕЛЬ КНОПОК
        # ========================================================

        toolbar = tk.Frame(
            self.root,
            bg=self.PANEL,
            highlightbackground=self.BORDER,
            highlightthickness=1
        )

        toolbar.pack(
            fill="x",
            padx=22,
            pady=(5, 12)
        )

        toolbar_inner = tk.Frame(
            toolbar,
            bg=self.PANEL
        )

        self.toolbar_inner = toolbar_inner
        self._panel_frames = [toolbar, toolbar_inner]

        toolbar_inner.pack(
            fill="x",
            padx=10,
            pady=9
        )

        # --------------------------------------------------------
        # PDF
        # --------------------------------------------------------

        self.select_button = ttk.Button(
            toolbar_inner,
            text=self.t("select"),
            command=self.select_pdfs,
            style="Accent.TButton"
        )

        self.select_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # Excel
        # --------------------------------------------------------

        self.excel_button = ttk.Button(
            toolbar_inner,
            text=self.t("excel"),
            command=self.save_excel,
            state="disabled",
            style="Dark.TButton"
        )

        self.excel_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # CSV
        # --------------------------------------------------------

        self.csv_button = ttk.Button(
            toolbar_inner,
            text=self.t("csv"),
            command=self.save_csv,
            state="disabled",
            style="Dark.TButton"
        )

        self.csv_button.pack(
            side="left",
            padx=3
        )

        self.print_button = ttk.Button(
            toolbar_inner,
            text=self.t("zpl"),
            command=self.print_zpl,
            state="disabled",
            style="Dark.TButton"
        )

        self.print_button.pack(
            side="left",
            padx=3
        )

        separator = tk.Frame(
            toolbar_inner,
            width=1,
            bg=self.BORDER
        )

        separator.pack(
            side="left",
            fill="y",
            padx=9
        )

        # --------------------------------------------------------
        # Копировать
        # --------------------------------------------------------

        self.copy_button = ttk.Button(
            toolbar_inner,
            text=self.t("copy"),
            command=self.copy_selected_code,
            state="disabled",
            style="Dark.TButton"
        )

        self.copy_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # Подробности
        # --------------------------------------------------------

        self.details_button = ttk.Button(
            toolbar_inner,
            text=self.t("details"),
            command=self.show_code_details,
            state="disabled",
            style="Dark.TButton"
        )

        self.details_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # Очистить
        # --------------------------------------------------------

        self.clear_button = ttk.Button(
            toolbar_inner,
            text=self.t("clear"),
            command=self.clear_results,
            style="Dark.TButton"
        )

        self.clear_button.pack(
            side="right",
            padx=3
        )

        # --------------------------------------------------------
        # Остановить
        # --------------------------------------------------------

        self.stop_button = ttk.Button(
            toolbar_inner,
            text=self.t("stop"),
            command=self.stop_scan,
            state="disabled",
            style="Danger.TButton"
        )

        self.stop_button.pack(
            side="right",
            padx=3
        )

        # --------------------------------------------------------
        # Настройки
        # --------------------------------------------------------

        self.settings_button = ttk.Button(
            toolbar_inner,
            text="⚙  " + self.t("settings"),
            command=self._open_settings,
            style="Dark.TButton"
        )
        self.settings_button.pack(side="right", padx=3)

        # ========================================================
        # КАРТОЧКИ СТАТИСТИКИ
        # ========================================================

        stats_frame = tk.Frame(
            self.root,
            bg=self.BG
        )

        self._bg_frames.append(stats_frame)
        self._panel_frames = getattr(self, "_panel_frames", [])

        stats_frame.pack(
            fill="x",
            padx=22,
            pady=(0, 12)
        )

        pdf_card = self.create_stat_card(
            stats_frame,
            self.t("pdf_stat"),
            "0"
        )

        pdf_card.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 6)
        )

        self.pdf_card_title = pdf_card.title_label
        self.pdf_count_value = pdf_card.value_label

        code_card = self.create_stat_card(
            stats_frame,
            self.t("code_stat"),
            "0"
        )

        code_card.pack(
            side="left",
            fill="x",
            expand=True,
            padx=6
        )

        self.code_card_title = code_card.title_label
        self.code_count_value = code_card.value_label

        status_card = self.create_stat_card(
            stats_frame,
            self.t("status_stat"),
            self.t("ready")
        )

        status_card.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(6, 0)
        )

        self.status_card_title = status_card.title_label
        self.status_value = status_card.value_label

        # ========================================================
        # ОСНОВНАЯ ОБЛАСТЬ
        # ========================================================

        main_frame = tk.Frame(
            self.root,
            bg=self.BG
        )

        self._bg_frames.append(main_frame)

        main_frame.pack(
            fill="both",
            expand=True,
            padx=22,
            pady=(0, 8)
        )

        # ========================================================
        # ЛЕВАЯ ПАНЕЛЬ
        # ========================================================

        left_frame = ttk.LabelFrame(
            main_frame,
            text=f"  {self.t('pdfs')}  ",
            style="Dark.TLabelframe"
        )

        self.file_frame = left_frame

        left_frame.pack(
            side="left",
            fill="both",
            pady=0,
            padx=(0, 6)
        )
        left_frame.configure(width=390)
        left_frame.pack_propagate(False)

        self.file_tree = ttk.Treeview(
            left_frame,
            columns=("count",),
            show="tree headings",
            height=20,
            style="Dark.Treeview"
        )

        self.file_tree.heading("#0", text=self.t("file"))
        self.file_tree.heading("count", text=self.t("codes"))
        self.file_tree.column("#0", width=315, minwidth=180)
        self.file_tree.column("count", width=70, anchor="center")

        self.file_tree.pack(
            fill="both",
            expand=True,
            padx=8,
            pady=8
        )

        # Вертикальная прокрутка PDF-панели убрана по просьбе пользователя.
        # Колесо мыши продолжает работать для длинных списков.
        self.file_tree.bind("<MouseWheel>", self._file_tree_mousewheel)
        self.file_tree.bind("<Button-4>", self._file_tree_mousewheel)
        self.file_tree.bind("<Button-5>", self._file_tree_mousewheel)

        # Изменение ширины панели курсором за правую границу.
        self._file_panel_resizing = False
        self._file_panel_resize_start_x = 0
        self._file_panel_resize_start_width = 390
        for widget in (left_frame, self.file_tree):
            widget.bind("<Motion>", self._file_panel_motion_cursor, add="+")
            widget.bind("<Button-1>", self._file_panel_press, add="+")
            widget.bind("<B1-Motion>", self._file_panel_motion, add="+")
            widget.bind("<ButtonRelease-1>", self._file_panel_release, add="+")

        self.file_tree.bind("<<TreeviewSelect>>", self.on_file_select)

        # ========================================================
        # ПРАВАЯ ПАНЕЛЬ
        # ========================================================

        right_frame = ttk.LabelFrame(
            main_frame,
            text=f"  {self.t('marking')}  ",
            style="Dark.TLabelframe"
        )

        self.code_frame = right_frame

        right_frame.pack(
            side="left",
            fill="both",
            expand=True,
            pady=0,
            padx=(6, 0)
        )

        columns = (
            "number",
            "page",
            "code"
        )

        self.code_tree = ttk.Treeview(
            right_frame,
            columns=columns,
            show="headings",
            height=20,
            selectmode="extended",
            style="Dark.Treeview"
        )

        self.code_tree.heading(
            "number",
            text=self.t("number")
        )

        self.code_tree.heading(
            "page",
            text=self.t("page")
        )

        self.code_tree.heading(
            "code",
            text=self.t("full")
        )

        self.code_tree.column(
            "number",
            width=55,
            anchor="center"
        )

        self.code_tree.column(
            "page",
            width=85,
            anchor="center"
        )

        self.code_tree.column(
            "code",
            width=750,
            minwidth=400
        )

        code_scroll_y = ttk.Scrollbar(
            right_frame,
            orient="vertical",
            command=self.code_tree.yview,
            style="Dark.Vertical.TScrollbar"
        )

        code_scroll_x = ttk.Scrollbar(
            right_frame,
            orient="horizontal",
            command=self.code_tree.xview,
            style="Dark.Horizontal.TScrollbar"
        )

        self.code_tree.configure(
            yscrollcommand=code_scroll_y.set,
            xscrollcommand=code_scroll_x.set
        )

        self.code_tree.pack(
            side="top",
            fill="both",
            expand=True,
            padx=(8, 0),
            pady=(8, 0)
        )

        code_scroll_y.pack(
            side="right",
            fill="y",
            padx=(0, 8),
            pady=(8, 0)
        )

        code_scroll_x.pack(
            side="bottom",
            fill="x",
            padx=8,
            pady=(0, 8)
        )

        self.code_tree.bind(
            "<Double-1>",
            self.copy_selected_code
        )

        self.code_tree.bind(
            "<<TreeviewSelect>>",
            self.on_code_select
        )

        # ========================================================
        # НИЖНЯЯ ПАНЕЛЬ
        # ========================================================

        bottom_frame = tk.Frame(
            self.root,
            bg=self.BG
        )

        bottom_frame.pack(
            fill="x",
            padx=22,
            pady=(0, 15)
        )

        self.progress = ttk.Progressbar(
            bottom_frame,
            mode="determinate",
            style="Dark.Horizontal.TProgressbar"
        )

        self.progress.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 12)
        )

        self.status_label = tk.Label(
            bottom_frame,
            text=self.t("ready"),
            bg=self.BG,
            fg=self.TEXT_SECONDARY,
            font=("Segoe UI", 9)
        )

        self.status_label.pack(
            side="right"
        )

    # ============================================================
    # КАРТОЧКА
    # ============================================================

    def create_stat_card(
        self,
        parent,
        title,
        value
    ):

        card = tk.Frame(
            parent,
            bg=self.PANEL,
            highlightbackground=self.BORDER,
            highlightthickness=1,
            height=72
        )

        card.pack_propagate(
            False
        )

        title_label = tk.Label(
            card,
            text=title,
            bg=self.PANEL,
            fg=self.TEXT_SECONDARY,
            font=("Segoe UI", 8, "bold")
        )
        title_label.pack(anchor="w", padx=14, pady=(10, 0))

        value_label = tk.Label(
            card,
            text=value,
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 16, "bold")
        )

        card.title_label = title_label
        value_label.pack(
            anchor="w",
            padx=14,
            pady=(0, 5)
        )

        card.value_label = value_label
        self._panel_frames = getattr(self, "_panel_frames", []) + [card]
        self._panel_labels = getattr(self, "_panel_labels", []) + [title_label, value_label]
        value_label._is_value = True

        return card

    # ============================================================
    # ВЫБОР PDF
    # ============================================================

    def select_pdfs(self):

        if self.scanning:
            return

        files = filedialog.askopenfilenames(
            title=self.t("select_files"),
            filetypes=[
                ("PDF-файлы", "*.pdf"),
                (self.t("all_files"), "*.*")
            ]
        )

        if not files:
            return

        # --------------------------------------------------------
        # ВАЖНО:
        # Сначала сохраняем выбранные файлы,
        # затем очищаем результаты.
        # --------------------------------------------------------

        self.pdf_files = list(files)

        self.clear_results(
            keep_files=True
        )

        self.pdf_count_value.config(
            text=str(len(self.pdf_files))
        )

        self.update_status(
            self.t("selected", count=len(self.pdf_files))
        )

        self.start_scan()

    # ============================================================
    # ЗАПУСК
    # ============================================================

    def start_scan(self):

        if not self.pdf_files:
            return

        self.scanning = True
        self.stop_requested = False
        self.scan_errors.clear()
        self.file_stats = {
            pdf_path: 0
            for pdf_path in self.pdf_files
        }

        self.select_button.config(
            state="disabled"
        )

        self.stop_button.config(
            state="normal"
        )

        self.excel_button.config(
            state="disabled"
        )

        self.csv_button.config(
            state="disabled"
        )

        self.header_status.config(
            text="● "+self.t("scanning"),
            fg=self.ACCENT
        )

        self.status_value.config(
            text=self.t("scanning")
        )

        thread = threading.Thread(
            target=self.scan_worker,
            daemon=True
        )

        thread.start()

    # ============================================================
    # ОСТАНОВКА
    # ============================================================

    def stop_scan(self):

        if self.scanning:

            self.stop_requested = True

            self.update_status(
                self.t("stopping")
            )

            self.header_status.config(
                text="● "+self.t("stopping"),
                fg=self.DANGER
            )

    # ============================================================
    # СКАНИРОВАНИЕ
    # ============================================================

    def scan_worker(self):

        total_pages = 0

        for pdf_path in self.pdf_files:

            try:

                document = fitz.open(
                    pdf_path
                )

                total_pages += len(
                    document
                )

                document.close()

            except Exception as error:
                self.root.after(
                    0,
                    self.add_scan_error,
                    pdf_path,
                    "Не удалось прочитать количество страниц",
                    error
                )

        if total_pages == 0:

            self.root.after(
                0,
                self.scan_finished
            )

            return

        processed_pages = 0

        for pdf_path in self.pdf_files:

            if self.stop_requested:
                break

            pdf_name = os.path.basename(pdf_path)

            try:

                document = fitz.open(
                    pdf_path
                )

            except Exception as error:

                self.root.after(
                    0,
                    self.add_scan_error,
                    pdf_path,
                    "Не удалось открыть PDF",
                    error
                )

                continue

            for page_index in range(
                len(document)
            ):

                if self.stop_requested:
                    break

                page = document[
                    page_index
                ]

                page_number = page_index + 1

                self.root.after(
                    0,
                    self.update_status,
                    f"{self.t('scanning')} {pdf_name} — {self.t('page').lower()} {page_number}"
                )

                # ------------------------------------------------
                # ИЗОБРАЖЕНИЯ ВНУТРИ PDF
                # ------------------------------------------------

                images = page.get_images(
                    full=True
                )

                for image_info in images:

                    if self.stop_requested:
                        break

                    xref = image_info[0]

                    try:

                        image_data = document.extract_image(
                            xref
                        )

                        image_bytes = image_data[
                            "image"
                        ]

                        image_array = np.frombuffer(
                            image_bytes,
                            dtype=np.uint8
                        )

                        image = cv2.imdecode(
                            image_array,
                            cv2.IMREAD_COLOR
                        )

                        if image is not None:

                            self.process_image(
                                image,
                                pdf_path,
                                page_number
                            )

                    except Exception as error:
                        self.root.after(
                            0,
                            self.add_scan_error,
                            pdf_path,
                            "Не удалось извлечь изображение",
                            error
                        )

                # ------------------------------------------------
                # РЕНДЕР
                # ------------------------------------------------

                for scale in [
                    3.0,
                    4.0,
                    5.0
                ]:

                    if self.stop_requested:
                        break

                    matrix = fitz.Matrix(
                        scale,
                        scale
                    )

                    try:

                        pix = page.get_pixmap(
                            matrix=matrix,
                            alpha=False
                        )

                        image_array = np.frombuffer(
                            pix.samples,
                            dtype=np.uint8
                        )

                        image = image_array.reshape(
                            pix.height,
                            pix.width,
                            pix.n
                        )

                        if pix.n == 4:

                            image = cv2.cvtColor(
                                image,
                                cv2.COLOR_RGBA2BGR
                            )

                        else:

                            image = cv2.cvtColor(
                                image,
                                cv2.COLOR_RGB2BGR
                            )

                        self.process_image(
                            image,
                            pdf_path,
                            page_number
                        )

                    except Exception as error:
                        self.root.after(
                            0,
                            self.add_scan_error,
                            pdf_path,
                            f"Не удалось отрендерить страницу {page_number}",
                            error
                        )

                processed_pages += 1

                progress_value = (
                    processed_pages /
                    total_pages
                ) * 100

                self.root.after(
                    0,
                    self.progress.configure,
                    {"value": progress_value}
                )

            document.close()

        self.root.after(
            0,
            self.scan_finished
        )

    # ============================================================
    # ОБРАБОТКА ИЗОБРАЖЕНИЯ
    # ============================================================

    def process_image(
        self,
        image,
        pdf_path,
        page_number
    ):

        if self.stop_requested:
            return

        variants = [
            image
        ]

        gray = cv2.cvtColor(
            image,
            cv2.COLOR_BGR2GRAY
        )

        variants.append(
            gray
        )

        clahe = cv2.createCLAHE(
            clipLimit=2.0,
            tileGridSize=(8, 8)
        )

        enhanced = clahe.apply(
            gray
        )

        variants.append(
            enhanced
        )

        _, otsu = cv2.threshold(
            enhanced,
            0,
            255,
            cv2.THRESH_BINARY +
            cv2.THRESH_OTSU
        )

        variants.append(
            otsu
        )

        adaptive = cv2.adaptiveThreshold(
            enhanced,
            255,
            cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
            cv2.THRESH_BINARY,
            31,
            11
        )

        variants.append(
            adaptive
        )

        for variant in variants:

            if self.stop_requested:
                return

            try:

                decoded = zxingcpp.read_barcodes(
                    variant
                )

                for barcode in decoded:

                    text = barcode.text

                    if not text:
                        continue

                    code = self.clean_code(
                        text
                    )

                    if not self.is_valid_marking_code(
                        code
                    ):
                        continue

                    self.root.after(
                        0,
                        self.add_result,
                        pdf_path,
                        page_number,
                        code
                    )

            except Exception as error:
                self.root.after(
                    0,
                    self.add_scan_error,
                    pdf_path,
                    "Ошибка распознавания штрихкода",
                    error
                )

    # ============================================================
    # ОЧИСТКА КОДА
    # ============================================================

    def clean_code(
        self,
        code
    ):

        if not code:
            return ""

        code = str(code).strip()

        # --------------------------------------------------------
        # НОРМАЛИЗАЦИЯ GS1
        # --------------------------------------------------------
        # В готовом коде Честного знака разделитель между AI
        # должен быть настоящим ASCII Group Separator (29).
        #
        # Если декодер вернул текстовую последовательность
        # "\x1d", заменяем её на настоящий символ chr(29).
        # --------------------------------------------------------

        code = code.replace(
            r"\x1d",
            chr(29)
        )

        code = code.replace(
            r"\u001d",
            chr(29)
        )

        # --------------------------------------------------------
        # УБИРАЕМ СКОБКИ AI
        # --------------------------------------------------------
        # (01)046... -> 01046...
        # (21)ABC... -> 21ABC...
        # (91)EE12   -> GS91EE12
        # (92)XYZ    -> GS92XYZ
        #
        # Важно: настоящий GS внутри уже отсканированного кода
        # не изменяется.
        # --------------------------------------------------------

        if code.startswith("(01)"):
            code = "01" + code[4:]

        code = code.replace(
            "(21)",
            "21",
            1
        )

        code = code.replace(
            "(91)",
            chr(29) + "91",
            1
        )

        code = code.replace(
            "(92)",
            chr(29) + "92",
            1
        )

        return code

    # ============================================================
    # ПРОВЕРКА
    # ============================================================

    def is_valid_marking_code(
        self,
        code
    ):

        if not code:
            return False

        normalized = code

        if not normalized.startswith("01"):
            return False

        gtin = normalized[2:16]

        return self.is_valid_gtin(gtin)

    def is_valid_gtin(self, gtin):

        if len(gtin) != 14 or not gtin.isdigit():
            return False

        check_sum = sum(
            int(digit) * (3 if index % 2 == 0 else 1)
            for index, digit in enumerate(reversed(gtin[:-1]))
        )
        check_digit = (10 - check_sum % 10) % 10

        return check_digit == int(gtin[-1])

    # ============================================================
    # ДОБАВЛЕНИЕ РЕЗУЛЬТАТА
    # ============================================================

    def add_result(
        self,
        pdf_path,
        page_number,
        code
    ):

        if code in self.seen_codes:
            return

        self.seen_codes.add(code)

        parsed = self.parse_marking_code(
            code
        )

        result = {
            "file": os.path.basename(pdf_path),
            "file_path": pdf_path,
            "page": page_number,
            "code": code,
            "gtin": parsed["gtin"],
            "serial": parsed["serial"],
            "crypto_91": parsed["crypto_91"],
            "crypto_92": parsed["crypto_92"],
            "crypto": parsed["crypto"]
        }

        self.results.append(
            result
        )

        self.file_stats[pdf_path] = (
            self.file_stats.get(
                pdf_path,
                0
            ) + 1
        )

        self.refresh_file_tree()
        self.update_code_count()

    # ============================================================
    # РАЗБОР КОДА
    # ============================================================

    def parse_marking_code(
        self,
        code
    ):

        result = {
            "gtin": "",
            "serial": "",
            "crypto_91": "",
            "crypto_92": "",
            "crypto": ""
        }

        if not code:
            return result

        raw = code

        if "(01)" in raw:

            raw = raw.replace(
                "(01)",
                "01"
            )

            raw = raw.replace(
                "(21)",
                "\x1d21"
            )

            raw = raw.replace(
                "(91)",
                "\x1d91"
            )

            raw = raw.replace(
                "(92)",
                "\x1d92"
            )

        # --------------------------------------------------------
        # GTIN
        # --------------------------------------------------------

        pos_01 = raw.find(
            "01"
        )

        if (
            pos_01 >= 0
            and len(raw) >= pos_01 + 16
        ):

            gtin = raw[
                pos_01 + 2:
                pos_01 + 16
            ]

            if gtin.isdigit():

                result["gtin"] = gtin

        # --------------------------------------------------------
        # AI 21
        # --------------------------------------------------------

        pos_21 = raw.find(
            "\x1d21"
        )

        if pos_21 >= 0:

            serial_start = (
                pos_21 + 3
            )

            serial_end = len(
                raw
            )

            for marker in [
                "\x1d91",
                "\x1d92",
                "\x1d93",
                "\x1d94",
                "\x1d95",
                "\x1d96",
                "\x1d97",
                "\x1d98",
                "\x1d99"
            ]:

                found = raw.find(
                    marker,
                    serial_start
                )

                if found >= 0:

                    serial_end = min(
                        serial_end,
                        found
                    )

            result["serial"] = raw[
                serial_start:
                serial_end
            ]

        else:

            if pos_01 >= 0:

                gtin_end = (
                    pos_01 + 16
                )

                pos_21_plain = raw.find(
                    "21",
                    gtin_end
                )

                if pos_21_plain >= 0:

                    serial_start = (
                        pos_21_plain + 2
                    )

                    serial_end = len(
                        raw
                    )

                    for marker in [
                        "91",
                        "92"
                    ]:

                        found = raw.find(
                            marker,
                            serial_start
                        )

                        if found >= 0:

                            serial_end = min(
                                serial_end,
                                found
                            )

                    result["serial"] = raw[
                        serial_start:
                        serial_end
                    ]

        # --------------------------------------------------------
        # AI 91
        # --------------------------------------------------------

        pos_91 = raw.find(
            "\x1d91"
        )

        if pos_91 >= 0:

            start = (
                pos_91 + 3
            )

            end = raw.find(
                "\x1d92",
                start
            )

            if end < 0:
                end = len(raw)

            result["crypto_91"] = raw[
                start:end
            ]

        else:

            pos_91_plain = raw.find(
                "91"
            )

            if pos_91_plain >= 0:

                start = (
                    pos_91_plain + 2
                )

                end = raw.find(
                    "92",
                    start
                )

                if end < 0:
                    end = len(raw)

                result["crypto_91"] = raw[
                    start:end
                ]

        # --------------------------------------------------------
        # AI 92
        # --------------------------------------------------------

        pos_92 = raw.find(
            "\x1d92"
        )

        if pos_92 >= 0:

            start = (
                pos_92 + 3
            )

            result["crypto_92"] = raw[
                start:
            ]

        else:

            pos_92_plain = raw.find(
                "92"
            )

            if pos_92_plain >= 0:

                start = (
                    pos_92_plain + 2
                )

                result["crypto_92"] = raw[
                    start:
                ]

        # --------------------------------------------------------
        # КРИПТОХВОСТ
        # --------------------------------------------------------

        crypto_parts = []

        if result["crypto_91"]:

            crypto_parts.append(
                "(91)" +
                result["crypto_91"]
            )

        if result["crypto_92"]:

            crypto_parts.append(
                "(92)" +
                result["crypto_92"]
            )

        result["crypto"] = "".join(
            crypto_parts
        )

        return result

    # ============================================================
    # ЛЕВОЕ ДЕРЕВО
    # ============================================================

    def _file_tree_mousewheel(self, event):
        if getattr(event, "num", None) == 4:
            delta = -1
        elif getattr(event, "num", None) == 5:
            delta = 1
        else:
            delta = -1 * int(event.delta / 120) if event.delta else 0
        if delta:
            self.file_tree.yview_scroll(delta, "units")
        return "break"

    def _file_panel_near_edge(self, event):
        try:
            panel_right = self.file_frame.winfo_rootx() + self.file_frame.winfo_width()
            return abs(event.x_root - panel_right) <= 8
        except Exception:
            return False

    def _file_panel_press(self, event):
        if not self._file_panel_near_edge(event):
            return
        self._file_panel_resizing = True
        self._file_panel_resize_start_x = event.x_root
        self._file_panel_resize_start_width = self.file_frame.winfo_width()
        self.root.configure(cursor="sb_h_double_arrow")

    def _file_panel_motion(self, event):
        if not self._file_panel_resizing:
            if self._file_panel_near_edge(event):
                self.root.configure(cursor="sb_h_double_arrow")
            else:
                self.root.configure(cursor="")
            return
        delta = event.x_root - self._file_panel_resize_start_x
        new_width = max(240, min(850, self._file_panel_resize_start_width + delta))
        self.file_frame.configure(width=new_width)
        self.file_frame.pack_propagate(False)

    def _file_panel_motion_cursor(self, event):
        if self._file_panel_resizing or self._file_panel_near_edge(event):
            self.root.configure(cursor="sb_h_double_arrow")
        else:
            self.root.configure(cursor="")

    def _file_panel_release(self, event):
        self._file_panel_resizing = False
        self.root.configure(cursor="")

    def refresh_file_tree(self):

        for item in self.file_tree.get_children():

            self.file_tree.delete(
                item
            )

        for index, pdf_path in enumerate(self.pdf_files, start=1):

            name = os.path.basename(
                pdf_path
            )

            count = self.file_stats.get(pdf_path, 0)

            self.file_tree.insert(
                "",
                "end",
                iid=pdf_path,
                text=f"{index}. {name}",
                values=(count,)
            )

    # ============================================================
    # ВЫБОР PDF
    # ============================================================

    def on_file_select(
        self,
        event=None
    ):

        selected = (
            self.file_tree.selection()
        )

        if not selected:
            return

        pdf_path = selected[0]

        self.code_tree.delete(
            *self.code_tree.get_children()
        )

        counter = 1

        for item in self.results:

            if item["file_path"] != pdf_path:
                continue

            self.code_tree.insert(
                "",
                "end",
                iid=str(id(item)),
                values=(
                    counter,
                    item["page"],
                    item["code"]
                )
            )

            counter += 1

    # ============================================================
    # ВЫБОР КОДА
    # ============================================================

    def on_code_select(
        self,
        event=None
    ):

        selected = (
            self.code_tree.selection()
        )

        if selected:

            self.copy_button.config(
                state="normal"
            )

            self.details_button.config(
                state="normal"
            )

        else:

            self.copy_button.config(
                state="disabled"
            )

            self.details_button.config(
                state="disabled"
            )

    # ============================================================
    # ПОЛУЧЕНИЕ КОДА
    # ============================================================

    def get_selected_result(self):

        selected = (
            self.code_tree.selection()
        )

        if not selected:
            return None

        selected_id = selected[0]

        for item in self.results:

            if str(id(item)) == selected_id:

                return item

        return None

    # ============================================================
    # КОПИРОВАНИЕ
    # ============================================================

    def copy_selected_code(
        self,
        event=None
    ):

        item = self.get_selected_result()

        if not item:
            return

        self.root.clipboard_clear()

        self.root.clipboard_append(
            item["code"]
        )

        self.root.update()

        self.update_status(
            self.t("copied_full")
        )

    # ============================================================
    # ПОДРОБНОСТИ
    # ============================================================

    def show_code_details(self):

        item = self.get_selected_result()
        if not item:
            return

        window = tk.Toplevel(self.root)
        window.title(self.t("details_title"))
        window.geometry("900x650")
        window.minsize(700, 500)
        window.configure(bg=self.BG)
        window.transient(self.root)

        header = tk.Frame(window, bg=self.BG)
        header.pack(fill="x", padx=25, pady=(20, 12))

        title_label = tk.Label(
            header, text=self.t("details_title"), bg=self.BG, fg=self.TEXT,
            font=("Segoe UI", 18, "bold")
        )
        title_label.pack(anchor="w")

        subtitle_label = tk.Label(
            header, text=self.t("details_subtitle"), bg=self.BG,
            fg=self.TEXT_SECONDARY, font=("Segoe UI", 9)
        )
        subtitle_label.pack(anchor="w")

        frame = tk.Frame(
            window, bg=self.PANEL,
            highlightbackground=self.BORDER, highlightthickness=1
        )
        frame.pack(fill="both", expand=True, padx=25, pady=(0, 20))

        content = tk.Frame(frame, bg=self.PANEL)
        content.pack(fill="both", expand=True, padx=20, pady=20)

        fields = []
        label, entry = self.create_detail_field(content, "GTIN", item["gtin"])
        fields.append((label, entry, "gtin"))
        label, entry = self.create_detail_field(content, self.t("serial_ai21"), item["serial"])
        fields.append((label, entry, "serial_ai21"))
        label, entry = self.create_detail_field(content, self.t("crypto91"), item["crypto_91"])
        fields.append((label, entry, "crypto91"))
        label, entry = self.create_detail_field(content, self.t("crypto92"), item["crypto_92"])
        fields.append((label, entry, "crypto92"))

        full_label = tk.Label(
            content, text=self.t("full"), bg=self.PANEL, fg=self.TEXT,
            font=("Segoe UI", 10, "bold")
        )
        full_label.pack(anchor="w", pady=(8, 5))

        text_frame = tk.Frame(
            content, bg=self.PANEL_2,
            highlightbackground=self.BORDER, highlightthickness=1
        )
        text_frame.pack(fill="both", expand=True)

        full_code_text = tk.Text(
            text_frame, wrap="none", bg=self.PANEL_2, fg=self.TEXT,
            insertbackground=self.TEXT, selectbackground=self.SELECTED,
            selectforeground="white", relief="flat", borderwidth=0,
            font=("Consolas", 10), padx=10, pady=10
        )
        full_code_text.pack(side="left", fill="both", expand=True)
        full_code_text.insert("1.0", item["code"])
        full_code_text.config(state="disabled")

        scroll = ttk.Scrollbar(
            text_frame, orient="vertical", command=full_code_text.yview,
            style="Dark.Vertical.TScrollbar"
        )
        scroll.pack(side="right", fill="y")
        full_code_text.configure(yscrollcommand=scroll.set)

        buttons = tk.Frame(content, bg=self.PANEL)
        buttons.pack(fill="x", pady=(15, 0))

        def copy_value(value, message_key):
            self.root.clipboard_clear()
            self.root.clipboard_append(value or "")
            self.root.update()
            self.update_status(self.t(message_key))

        copy_gtin = ttk.Button(
            buttons, text=self.t("copy_gtin"),
            command=lambda: copy_value(item["gtin"], "copied_gtin"),
            style="Dark.TButton"
        )
        copy_gtin.pack(side="left", padx=(0, 5))

        copy_serial = ttk.Button(
            buttons, text=self.t("copy_serial"),
            command=lambda: copy_value(item["serial"], "copied_serial"),
            style="Dark.TButton"
        )
        copy_serial.pack(side="left", padx=5)

        copy_crypto = ttk.Button(
            buttons, text=self.t("copy_crypto"),
            command=lambda: copy_value(item["crypto"], "copied_crypto"),
            style="Dark.TButton"
        )
        copy_crypto.pack(side="left", padx=5)

        close_button = ttk.Button(
            buttons, text=self.t("close"), command=window.destroy,
            style="Accent.TButton"
        )
        close_button.pack(side="right")

        refs = {
            "window": window, "header": header, "title": title_label,
            "subtitle": subtitle_label, "frame": frame, "content": content,
            "full_label": full_label, "text_frame": text_frame,
            "full_text": full_code_text, "buttons_frame": buttons,
            "copy_gtin": copy_gtin, "copy_serial": copy_serial,
            "copy_crypto": copy_crypto, "close": close_button,
            "fields": fields
        }
        self._detail_windows.append(refs)

        def on_close():
            self._detail_windows = [r for r in self._detail_windows if r is not refs]
            try:
                window.destroy()
            except Exception:
                pass

        close_button.configure(command=on_close)
        window.protocol("WM_DELETE_WINDOW", on_close)

    # ============================================================
    # ПОЛЕ ПОДРОБНОСТЕЙ
    # ============================================================

    def create_detail_field(self, parent, title, value):

        label = tk.Label(
            parent,
            text=title,
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 10, "bold")
        )
        label.pack(anchor="w", pady=(0, 4))

        entry = tk.Entry(
            parent,
            bg=self.PANEL_2,
            fg=self.TEXT,
            insertbackground=self.TEXT,
            readonlybackground=self.PANEL_2,
            relief="flat",
            font=("Consolas", 10),
            highlightbackground=self.BORDER,
            highlightthickness=1
        )
        entry.pack(fill="x", pady=(0, 10), ipady=6)
        entry.insert(0, value or "")
        entry.config(state="readonly")
        return label, entry

    # ============================================================
    # СТАТИСТИКА
    # ============================================================

    def update_code_count(self):

        count = len(
            self.results
        )

        self.code_count_value.config(
            text=str(count)
        )

    # ============================================================
    # СТАТУС
    # ============================================================

    def update_status(
        self,
        text
    ):

        self.status_label.config(
            text=text
        )

        self.status_value.config(
            text=text[:22]
        )

    # ============================================================
    # ОШИБКА
    # ============================================================

    def show_error(
        self,
        title,
        message
    ):

        messagebox.showerror(title, message)

    # ============================================================
    # ЗАВЕРШЕНИЕ
    # ============================================================

    def scan_finished(self):

        self.scanning = False

        self.select_button.config(
            state="normal"
        )

        self.stop_button.config(
            state="disabled"
        )

        if self.results:

            self.excel_button.config(
                state="normal"
            )

            self.csv_button.config(
                state="normal"
            )

            self.print_button.config(
                state="normal"
            )

        self.progress["value"] = 100

        if self.stop_requested:

            self.update_status(
                self.t("stopped")
            )

            self.header_status.config(
                text="● "+self.t("stopped"),
                fg=self.DANGER
            )

            self.status_value.config(
                text=self.t("stopped")
            )

        else:

            error_suffix = (
                f" Ошибок: {len(self.scan_errors)}"
                if self.scan_errors else ""
            )

            self.update_status(
                self.t("found", count=len(self.results)) + ("." + error_suffix if error_suffix else "")
            )

            self.header_status.config(
                text="● "+self.t("done"),
                fg=self.SUCCESS
            )

            self.status_value.config(
                text=self.t("done")
            )

        self.refresh_file_tree()

    def add_scan_error(self, pdf_path, context, error):

        message = f"{os.path.basename(pdf_path)}: {context}: {error}"

        if message not in self.scan_errors:
            self.scan_errors.append(message)

    # ============================================================
    # ОЧИСТКА
    # ============================================================

    def clear_results(
        self,
        keep_files=False
    ):

        self.results.clear()

        self.file_stats.clear()

        self.seen_codes.clear()

        self.scan_errors.clear()

        self.file_tree.delete(
            *self.file_tree.get_children()
        )

        self.code_tree.delete(
            *self.code_tree.get_children()
        )

        self.code_count_value.config(
            text="0"
        )

        self.progress["value"] = 0

        self.copy_button.config(
            state="disabled"
        )

        self.details_button.config(
            state="disabled"
        )

        self.excel_button.config(
            state="disabled"
        )

        self.csv_button.config(
            state="disabled"
        )

        self.print_button.config(
            state="disabled"
        )

        if not keep_files:

            self.pdf_files = []

            self.pdf_count_value.config(
                text="0"
            )

        self.update_status(
            self.t("ready")
        )

        self.header_status.config(
            text="● "+self.t("ready"),
            fg=self.SUCCESS
        )

        self.status_value.config(
            text=self.t("ready")
        )

    # ============================================================
    # ПЕЧАТЬ ЭТИКЕТОК ZPL
    # ============================================================

    def get_selected_results(self):

        selected_ids = set(self.code_tree.selection())

        return [
            item
            for item in self.results
            if str(id(item)) in selected_ids
        ]

    def zpl_escape(self, value):

        escaped = []

        for byte in value.encode("utf-8"):

            char = chr(byte)

            if 32 <= byte <= 126 and char not in "^~\\":
                escaped.append(char)
            else:
                escaped.append(f"\\{byte:02X}")

        return "".join(escaped)

    def build_zpl_label(self, item, width_mm, height_mm, dots_per_mm):

        width = width_mm * dots_per_mm
        height = height_mm * dots_per_mm
        margin = 3 * dots_per_mm
        matrix_size = max(3, min(8, int(height / 65)))
        text_y = margin + matrix_size * 30 + (2 * dots_per_mm)

        code = self.zpl_escape(item["code"])
        gtin = self.zpl_escape(item["gtin"] or "—")
        serial = self.zpl_escape(item["serial"] or "—")

        return (
            "^XA\n"
            f"^PW{width}\n"
            f"^LL{height}\n"
            "^CI28\n"
            f"^FO{margin},{margin}^BXN,{matrix_size},200^FH\\^FD{code}^FS\n"
            f"^FO{margin},{text_y}^A0N,{2 * dots_per_mm},{2 * dots_per_mm}^FDGTIN: {gtin}^FS\n"
            f"^FO{margin},{text_y + 3 * dots_per_mm}^A0N,{2 * dots_per_mm},{2 * dots_per_mm}^FDSN: {serial}^FS\n"
            "^XZ\n"
        )

    def print_zpl(self):

        selected = self.get_selected_results()

        if selected:
            use_selected = messagebox.askyesno(self.t("print_title"), self.t("print_question", count=len(selected)))
            items = selected if use_selected else self.results
        else:
            items = self.results

        if not items:
            return

        width_mm = simpledialog.askinteger(
            self.t("label_size"),
            self.t("label_width"),
            parent=self.root,
            initialvalue=58,
            minvalue=20,
            maxvalue=200
        )

        if width_mm is None:
            return

        height_mm = simpledialog.askinteger(
            self.t("label_size"),
            self.t("label_height"),
            parent=self.root,
            initialvalue=40,
            minvalue=20,
            maxvalue=200
        )

        if height_mm is None:
            return

        zpl = "".join(
            self.build_zpl_label(item, width_mm, height_mm, 8)
            for item in items
        )

        try:
            import win32print

            printer_name = win32print.GetDefaultPrinter()

            if not messagebox.askyesno(
                self.t("confirm_print"), self.t("send_print", count=len(items), printer=printer_name)
            ):
                return

            printer = win32print.OpenPrinter(printer_name)

            try:
                job = win32print.StartDocPrinter(
                    printer,
                    1,
                    ("Chestny Znak labels", None, "RAW")
                )
                win32print.StartPagePrinter(printer)
                win32print.WritePrinter(printer, zpl.encode("utf-8"))
                win32print.EndPagePrinter(printer)
                win32print.EndDocPrinter(printer)
            finally:
                win32print.ClosePrinter(printer)

            self.update_status(self.t("sent", count=len(items)))

        except ImportError:
            self.save_zpl_file(zpl)

        except Exception as error:
            messagebox.showerror(self.t("print_error"), str(error))

    def save_zpl_file(self, zpl):

        filename = filedialog.asksaveasfilename(
            title=self.t("save_zpl"),
            defaultextension=".zpl",
            filetypes=[
                (self.t("zpl_file"), "*.zpl"),
                (self.t("all_files"), "*.*")
            ]
        )

        if not filename:
            return

        try:
            with open(filename, "w", encoding="utf-8", newline="") as file:
                file.write(zpl)

            messagebox.showinfo(self.t("zpl_saved"), self.t("zpl_fallback"))

        except Exception as error:
            messagebox.showerror(self.t("save_error"), str(error))

    # ============================================================
    # EXCEL
    # ============================================================

    def save_excel(self):

        if not self.results:
            return

        filename = filedialog.asksaveasfilename(
            title=self.t("save_excel"),
            defaultextension=".xlsx",
            filetypes=[
                (self.t("excel_file"), "*.xlsx"),
                (self.t("all_files"), "*.*")
            ]
        )

        if not filename:
            return

        try:

            workbook = Workbook()

            sheet = workbook.active

            sheet.title = self.t("marking")

            headers = [
                self.t("number"),
                self.t("file"),
                self.t("page"),
                "GTIN",
                self.t("serial"),
                self.t("crypto91"),
                self.t("crypto92"),
                self.t("crypto"),
                self.t("full")
            ]

            sheet.append(
                headers
            )

            for index, item in enumerate(
                self.results,
                start=1
            ):

                sheet.append([
                    index,
                    item["file"],
                    item["page"],
                    item["gtin"],
                    item["serial"],
                    item["crypto_91"],
                    item["crypto_92"],
                    item["crypto"],
                    item["code"]
                ])

            sheet.auto_filter.ref = (
                f"A1:I{sheet.max_row}"
            )

            sheet.freeze_panes = "A2"

            widths = {
                "A": 8,
                "B": 45,
                "C": 12,
                "D": 18,
                "E": 30,
                "F": 35,
                "G": 35,
                "H": 70,
                "I": 100
            }

            for column, width in widths.items():

                sheet.column_dimensions[
                    column
                ].width = width

            workbook.save(
                filename
            )

            messagebox.showinfo(self.t("done"), self.t("excel_saved", file=filename))

        except Exception as e:

            messagebox.showerror(
                self.t("excel_error"),
                str(e)
            )

    # ============================================================
    # CSV
    # ============================================================

    def save_csv(self):

        if not self.results:
            return

        filename = filedialog.asksaveasfilename(
            title=self.t("save_csv"),
            defaultextension=".csv",
            filetypes=[
                (self.t("csv_file"), "*.csv"),
                (self.t("all_files"), "*.*")
            ]
        )

        if not filename:
            return

        try:

            with open(
                filename,
                "w",
                newline="",
                encoding="utf-8-sig"
            ) as file:

                writer = csv.writer(
                    file,
                    delimiter=";",
                    quoting=csv.QUOTE_MINIMAL
                )

                writer.writerow([
                    self.t("number"),
                    self.t("file"),
                    self.t("page"),
                    "GTIN",
                    self.t("serial"),
                    self.t("crypto91"),
                    self.t("crypto92"),
                    self.t("crypto"),
                    self.t("full")
                ])

                for index, item in enumerate(
                    self.results,
                    start=1
                ):

                    writer.writerow([
                        index,
                        item["file"],
                        item["page"],
                        item["gtin"],
                        item["serial"],
                        item["crypto_91"],
                        item["crypto_92"],
                        item["crypto"],
                        item["code"]
                    ])

            messagebox.showinfo(self.t("done"), self.t("csv_saved", file=filename))

        except Exception as e:

            messagebox.showerror(
                self.t("csv_error"),
                str(e)
            )


# ================================================================
# ЗАПУСК
# ================================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = ChestnyZnakScanner(
        root
    )

    root.mainloop()
