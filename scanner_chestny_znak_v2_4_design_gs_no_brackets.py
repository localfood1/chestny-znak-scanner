import tkinter as tk
from tkinter import filedialog, messagebox, simpledialog, ttk
import threading
import os
import csv
import ctypes

import fitz
import cv2
import numpy as np
import zxingcpp
from openpyxl import Workbook


# ================================================================
# DPI ДЛЯ WINDOWS
# ================================================================

try:
    ctypes.windll.shcore.SetProcessDpiAwareness(1)
except Exception:
    try:
        ctypes.windll.user32.SetProcessDPIAware()
    except Exception:
        pass


class ChestnyZnakScanner:

    # ============================================================
    # ЦВЕТА
    # ============================================================

    BG = "#111318"
    PANEL = "#181B21"
    PANEL_2 = "#20242C"
    BORDER = "#2C313B"

    TEXT = "#F1F3F5"
    TEXT_SECONDARY = "#9BA3AF"

    ACCENT = "#4C8DFF"
    ACCENT_HOVER = "#6EA3FF"

    SUCCESS = "#35C98B"
    DANGER = "#E85D6A"

    SELECTED = "#263D63"

    # ============================================================
    # ИНИЦИАЛИЗАЦИЯ
    # ============================================================

    def __init__(self, root):

        self.root = root

        self.root.title(
            "Честный знак — Сканер маркировки"
        )

        self.root.geometry(
            "1250x750"
        )

        self.root.minsize(
            1000,
            600
        )

        self.root.configure(
            bg=self.BG
        )

        # --------------------------------------------------------
        # Данные
        # --------------------------------------------------------

        self.pdf_files = []
        self.results = []
        self.file_stats = {}
        self.seen_codes = set()
        self.scan_errors = []

        self.scanning = False
        self.stop_requested = False

        # --------------------------------------------------------
        # Интерфейс
        # --------------------------------------------------------

        self.setup_styles()
        self.build_interface()

    # ============================================================
    # СТИЛИ
    # ============================================================

    def setup_styles(self):

        style = ttk.Style()

        try:
            style.theme_use("clam")
        except Exception:
            pass

        # --------------------------------------------------------
        # Основные элементы
        # --------------------------------------------------------

        style.configure(
            ".",
            background=self.BG,
            foreground=self.TEXT,
            font=("Segoe UI", 10)
        )

        style.configure(
            "TFrame",
            background=self.BG
        )

        style.configure(
            "TLabel",
            background=self.BG,
            foreground=self.TEXT
        )

        # --------------------------------------------------------
        # Заголовок
        # --------------------------------------------------------

        style.configure(
            "Title.TLabel",
            background=self.BG,
            foreground=self.TEXT,
            font=("Segoe UI", 18, "bold")
        )

        style.configure(
            "Subtitle.TLabel",
            background=self.BG,
            foreground=self.TEXT_SECONDARY,
            font=("Segoe UI", 9)
        )

        # --------------------------------------------------------
        # LabelFrame
        # --------------------------------------------------------

        style.configure(
            "Dark.TLabelframe",
            background=self.PANEL,
            foreground=self.TEXT,
            bordercolor=self.BORDER,
            relief="solid",
            borderwidth=1
        )

        style.configure(
            "Dark.TLabelframe.Label",
            background=self.PANEL,
            foreground=self.TEXT,
            font=("Segoe UI", 10, "bold")
        )

        # --------------------------------------------------------
        # Обычные кнопки
        # --------------------------------------------------------

        style.configure(
            "Dark.TButton",
            background=self.PANEL_2,
            foreground=self.TEXT,
            bordercolor=self.BORDER,
            borderwidth=1,
            padding=(13, 8),
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Dark.TButton",
            background=[
                ("active", "#2A303A"),
                ("pressed", "#303743"),
                ("disabled", "#15171B")
            ],
            foreground=[
                ("disabled", "#555B66")
            ]
        )

        # --------------------------------------------------------
        # Главная кнопка
        # --------------------------------------------------------

        style.configure(
            "Accent.TButton",
            background=self.ACCENT,
            foreground="white",
            borderwidth=0,
            padding=(15, 9),
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Accent.TButton",
            background=[
                ("active", self.ACCENT_HOVER),
                ("pressed", self.ACCENT)
            ]
        )

        # --------------------------------------------------------
        # Кнопка остановки
        # --------------------------------------------------------

        style.configure(
            "Danger.TButton",
            background=self.DANGER,
            foreground="white",
            borderwidth=0,
            padding=(13, 8),
            font=("Segoe UI", 9, "bold")
        )

        # --------------------------------------------------------
        # Treeview
        # --------------------------------------------------------

        style.configure(
            "Dark.Treeview",
            background=self.PANEL,
            fieldbackground=self.PANEL,
            foreground=self.TEXT,
            bordercolor=self.BORDER,
            rowheight=34,
            font=("Segoe UI", 9)
        )

        style.map(
            "Dark.Treeview",
            background=[
                ("selected", self.SELECTED)
            ],
            foreground=[
                ("selected", "white")
            ]
        )

        style.configure(
            "Dark.Treeview.Heading",
            background=self.PANEL_2,
            foreground=self.TEXT_SECONDARY,
            bordercolor=self.BORDER,
            relief="flat",
            padding=(8, 9),
            font=("Segoe UI", 9, "bold")
        )

        style.map(
            "Dark.Treeview.Heading",
            background=[
                ("active", "#2A303A")
            ]
        )

        # --------------------------------------------------------
        # Scrollbars
        # --------------------------------------------------------

        style.configure(
            "Dark.Vertical.TScrollbar",
            background=self.PANEL_2,
            troughcolor=self.PANEL,
            bordercolor=self.PANEL,
            arrowcolor=self.TEXT_SECONDARY
        )

        style.configure(
            "Dark.Horizontal.TScrollbar",
            background=self.PANEL_2,
            troughcolor=self.PANEL,
            bordercolor=self.PANEL,
            arrowcolor=self.TEXT_SECONDARY
        )

        # --------------------------------------------------------
        # Progressbar
        # --------------------------------------------------------

        style.configure(
            "Dark.Horizontal.TProgressbar",
            troughcolor=self.PANEL_2,
            background=self.ACCENT,
            bordercolor=self.PANEL_2,
            lightcolor=self.ACCENT,
            darkcolor=self.ACCENT
        )

    # ============================================================
    # ИНТЕРФЕЙС
    # ============================================================

    def build_interface(self):

        # ========================================================
        # ШАПКА
        # ========================================================

        header = tk.Frame(
            self.root,
            bg=self.BG
        )

        header.pack(
            fill="x",
            padx=22,
            pady=(18, 8)
        )

        title_frame = tk.Frame(
            header,
            bg=self.BG
        )

        title_frame.pack(
            side="left"
        )

        ttk.Label(
            title_frame,
            text="Честный знак",
            style="Title.TLabel"
        ).pack(
            anchor="w"
        )

        ttk.Label(
            title_frame,
            text="Сканер кодов маркировки из PDF",
            style="Subtitle.TLabel"
        ).pack(
            anchor="w",
            pady=(2, 0)
        )

        self.header_status = tk.Label(
            header,
            text="● Готов",
            bg=self.BG,
            fg=self.SUCCESS,
            font=("Segoe UI", 10, "bold")
        )

        self.header_status.pack(
            side="right",
            pady=5
        )

        # ========================================================
        # ПАНЕЛЬ КНОПОК
        # ========================================================

        toolbar = tk.Frame(
            self.root,
            bg=self.PANEL,
            highlightbackground=self.BORDER,
            highlightthickness=1
        )

        toolbar.pack(
            fill="x",
            padx=22,
            pady=(5, 12)
        )

        toolbar_inner = tk.Frame(
            toolbar,
            bg=self.PANEL
        )

        toolbar_inner.pack(
            fill="x",
            padx=10,
            pady=9
        )

        # --------------------------------------------------------
        # PDF
        # --------------------------------------------------------

        self.select_button = ttk.Button(
            toolbar_inner,
            text="📄  Выбрать PDF",
            command=self.select_pdfs,
            style="Accent.TButton"
        )

        self.select_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # Excel
        # --------------------------------------------------------

        self.excel_button = ttk.Button(
            toolbar_inner,
            text="📊  Excel",
            command=self.save_excel,
            state="disabled",
            style="Dark.TButton"
        )

        self.excel_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # CSV
        # --------------------------------------------------------

        self.csv_button = ttk.Button(
            toolbar_inner,
            text="📋  CSV",
            command=self.save_csv,
            state="disabled",
            style="Dark.TButton"
        )

        self.csv_button.pack(
            side="left",
            padx=3
        )

        self.print_button = ttk.Button(
            toolbar_inner,
            text="🏷  Печать ZPL",
            command=self.print_zpl,
            state="disabled",
            style="Dark.TButton"
        )

        self.print_button.pack(
            side="left",
            padx=3
        )

        separator = tk.Frame(
            toolbar_inner,
            width=1,
            bg=self.BORDER
        )

        separator.pack(
            side="left",
            fill="y",
            padx=9
        )

        # --------------------------------------------------------
        # Копировать
        # --------------------------------------------------------

        self.copy_button = ttk.Button(
            toolbar_inner,
            text="📋  Копировать",
            command=self.copy_selected_code,
            state="disabled",
            style="Dark.TButton"
        )

        self.copy_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # Подробности
        # --------------------------------------------------------

        self.details_button = ttk.Button(
            toolbar_inner,
            text="🔍  Подробности",
            command=self.show_code_details,
            state="disabled",
            style="Dark.TButton"
        )

        self.details_button.pack(
            side="left",
            padx=3
        )

        # --------------------------------------------------------
        # Очистить
        # --------------------------------------------------------

        self.clear_button = ttk.Button(
            toolbar_inner,
            text="🗑  Очистить",
            command=self.clear_results,
            style="Dark.TButton"
        )

        self.clear_button.pack(
            side="right",
            padx=3
        )

        # --------------------------------------------------------
        # Остановить
        # --------------------------------------------------------

        self.stop_button = ttk.Button(
            toolbar_inner,
            text="⛔  Остановить",
            command=self.stop_scan,
            state="disabled",
            style="Danger.TButton"
        )

        self.stop_button.pack(
            side="right",
            padx=3
        )

        # ========================================================
        # КАРТОЧКИ СТАТИСТИКИ
        # ========================================================

        stats_frame = tk.Frame(
            self.root,
            bg=self.BG
        )

        stats_frame.pack(
            fill="x",
            padx=22,
            pady=(0, 12)
        )

        pdf_card = self.create_stat_card(
            stats_frame,
            "PDF-ФАЙЛОВ",
            "0"
        )

        pdf_card.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 6)
        )

        self.pdf_count_value = pdf_card.value_label

        code_card = self.create_stat_card(
            stats_frame,
            "УНИКАЛЬНЫХ КОДОВ",
            "0"
        )

        code_card.pack(
            side="left",
            fill="x",
            expand=True,
            padx=6
        )

        self.code_count_value = code_card.value_label

        status_card = self.create_stat_card(
            stats_frame,
            "СТАТУС",
            "Готов"
        )

        status_card.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(6, 0)
        )

        self.status_value = status_card.value_label

        # ========================================================
        # ОСНОВНАЯ ОБЛАСТЬ
        # ========================================================

        main_frame = tk.Frame(
            self.root,
            bg=self.BG
        )

        main_frame.pack(
            fill="both",
            expand=True,
            padx=22,
            pady=(0, 8)
        )

        # ========================================================
        # ЛЕВАЯ ПАНЕЛЬ
        # ========================================================

        left_frame = ttk.LabelFrame(
            main_frame,
            text="  PDF-ФАЙЛЫ  ",
            style="Dark.TLabelframe"
        )

        left_frame.pack(
            side="left",
            fill="both",
            pady=0,
            padx=(0, 6)
        )

        self.file_tree = ttk.Treeview(
            left_frame,
            columns=("count",),
            show="tree headings",
            height=20,
            style="Dark.Treeview"
        )

        self.file_tree.heading(
            "#0",
            text="Файл"
        )

        self.file_tree.heading(
            "count",
            text="Кодов"
        )

        self.file_tree.column(
            "#0",
            width=280,
            minwidth=180
        )

        self.file_tree.column(
            "count",
            width=70,
            anchor="center"
        )

        file_scroll = ttk.Scrollbar(
            left_frame,
            orient="vertical",
            command=self.file_tree.yview,
            style="Dark.Vertical.TScrollbar"
        )

        self.file_tree.configure(
            yscrollcommand=file_scroll.set
        )

        self.file_tree.pack(
            side="left",
            fill="both",
            expand=True,
            padx=(8, 0),
            pady=8
        )

        file_scroll.pack(
            side="right",
            fill="y",
            padx=(0, 8),
            pady=8
        )

        self.file_tree.bind(
            "<<TreeviewSelect>>",
            self.on_file_select
        )

        # ========================================================
        # ПРАВАЯ ПАНЕЛЬ
        # ========================================================

        right_frame = ttk.LabelFrame(
            main_frame,
            text="  КОДЫ МАРКИРОВКИ  ",
            style="Dark.TLabelframe"
        )

        right_frame.pack(
            side="left",
            fill="both",
            expand=True,
            pady=0,
            padx=(6, 0)
        )

        columns = (
            "number",
            "page",
            "code"
        )

        self.code_tree = ttk.Treeview(
            right_frame,
            columns=columns,
            show="headings",
            height=20,
            selectmode="extended",
            style="Dark.Treeview"
        )

        self.code_tree.heading(
            "number",
            text="№"
        )

        self.code_tree.heading(
            "page",
            text="Страница"
        )

        self.code_tree.heading(
            "code",
            text="Полный код маркировки"
        )

        self.code_tree.column(
            "number",
            width=55,
            anchor="center"
        )

        self.code_tree.column(
            "page",
            width=85,
            anchor="center"
        )

        self.code_tree.column(
            "code",
            width=750,
            minwidth=400
        )

        code_scroll_y = ttk.Scrollbar(
            right_frame,
            orient="vertical",
            command=self.code_tree.yview,
            style="Dark.Vertical.TScrollbar"
        )

        code_scroll_x = ttk.Scrollbar(
            right_frame,
            orient="horizontal",
            command=self.code_tree.xview,
            style="Dark.Horizontal.TScrollbar"
        )

        self.code_tree.configure(
            yscrollcommand=code_scroll_y.set,
            xscrollcommand=code_scroll_x.set
        )

        self.code_tree.pack(
            side="top",
            fill="both",
            expand=True,
            padx=(8, 0),
            pady=(8, 0)
        )

        code_scroll_y.pack(
            side="right",
            fill="y",
            padx=(0, 8),
            pady=(8, 0)
        )

        code_scroll_x.pack(
            side="bottom",
            fill="x",
            padx=8,
            pady=(0, 8)
        )

        self.code_tree.bind(
            "<Double-1>",
            self.copy_selected_code
        )

        self.code_tree.bind(
            "<<TreeviewSelect>>",
            self.on_code_select
        )

        # ========================================================
        # НИЖНЯЯ ПАНЕЛЬ
        # ========================================================

        bottom_frame = tk.Frame(
            self.root,
            bg=self.BG
        )

        bottom_frame.pack(
            fill="x",
            padx=22,
            pady=(0, 15)
        )

        self.progress = ttk.Progressbar(
            bottom_frame,
            mode="determinate",
            style="Dark.Horizontal.TProgressbar"
        )

        self.progress.pack(
            side="left",
            fill="x",
            expand=True,
            padx=(0, 12)
        )

        self.status_label = tk.Label(
            bottom_frame,
            text="Готов",
            bg=self.BG,
            fg=self.TEXT_SECONDARY,
            font=("Segoe UI", 9)
        )

        self.status_label.pack(
            side="right"
        )

    # ============================================================
    # КАРТОЧКА
    # ============================================================

    def create_stat_card(
        self,
        parent,
        title,
        value
    ):

        card = tk.Frame(
            parent,
            bg=self.PANEL,
            highlightbackground=self.BORDER,
            highlightthickness=1,
            height=72
        )

        card.pack_propagate(
            False
        )

        tk.Label(
            card,
            text=title,
            bg=self.PANEL,
            fg=self.TEXT_SECONDARY,
            font=("Segoe UI", 8, "bold")
        ).pack(
            anchor="w",
            padx=14,
            pady=(10, 0)
        )

        value_label = tk.Label(
            card,
            text=value,
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 16, "bold")
        )

        value_label.pack(
            anchor="w",
            padx=14,
            pady=(0, 5)
        )

        card.value_label = value_label

        return card

    # ============================================================
    # ВЫБОР PDF
    # ============================================================

    def select_pdfs(self):

        if self.scanning:
            return

        files = filedialog.askopenfilenames(
            title="Выберите PDF-файлы",
            filetypes=[
                ("PDF-файлы", "*.pdf"),
                ("Все файлы", "*.*")
            ]
        )

        if not files:
            return

        # --------------------------------------------------------
        # ВАЖНО:
        # Сначала сохраняем выбранные файлы,
        # затем очищаем результаты.
        # --------------------------------------------------------

        self.pdf_files = list(files)

        self.clear_results(
            keep_files=True
        )

        self.pdf_count_value.config(
            text=str(len(self.pdf_files))
        )

        self.update_status(
            f"Выбрано PDF: {len(self.pdf_files)}"
        )

        self.start_scan()

    # ============================================================
    # ЗАПУСК
    # ============================================================

    def start_scan(self):

        if not self.pdf_files:
            return

        self.scanning = True
        self.stop_requested = False
        self.scan_errors.clear()
        self.file_stats = {
            pdf_path: 0
            for pdf_path in self.pdf_files
        }

        self.select_button.config(
            state="disabled"
        )

        self.stop_button.config(
            state="normal"
        )

        self.excel_button.config(
            state="disabled"
        )

        self.csv_button.config(
            state="disabled"
        )

        self.header_status.config(
            text="● Сканирование...",
            fg=self.ACCENT
        )

        self.status_value.config(
            text="Сканирование..."
        )

        thread = threading.Thread(
            target=self.scan_worker,
            daemon=True
        )

        thread.start()

    # ============================================================
    # ОСТАНОВКА
    # ============================================================

    def stop_scan(self):

        if self.scanning:

            self.stop_requested = True

            self.update_status(
                "Остановка..."
            )

            self.header_status.config(
                text="● Остановка...",
                fg=self.DANGER
            )

    # ============================================================
    # СКАНИРОВАНИЕ
    # ============================================================

    def scan_worker(self):

        total_pages = 0

        for pdf_path in self.pdf_files:

            try:

                document = fitz.open(
                    pdf_path
                )

                total_pages += len(
                    document
                )

                document.close()

            except Exception as error:
                self.root.after(
                    0,
                    self.add_scan_error,
                    pdf_path,
                    "Не удалось прочитать количество страниц",
                    error
                )

        if total_pages == 0:

            self.root.after(
                0,
                self.scan_finished
            )

            return

        processed_pages = 0

        for pdf_path in self.pdf_files:

            if self.stop_requested:
                break

            pdf_name = os.path.basename(pdf_path)

            try:

                document = fitz.open(
                    pdf_path
                )

            except Exception as error:

                self.root.after(
                    0,
                    self.add_scan_error,
                    pdf_path,
                    "Не удалось открыть PDF",
                    error
                )

                continue

            for page_index in range(
                len(document)
            ):

                if self.stop_requested:
                    break

                page = document[
                    page_index
                ]

                page_number = page_index + 1

                self.root.after(
                    0,
                    self.update_status,
                    f"Сканирование: {pdf_name} — страница {page_number}"
                )

                # ------------------------------------------------
                # ИЗОБРАЖЕНИЯ ВНУТРИ PDF
                # ------------------------------------------------

                images = page.get_images(
                    full=True
                )

                for image_info in images:

                    if self.stop_requested:
                        break

                    xref = image_info[0]

                    try:

                        image_data = document.extract_image(
                            xref
                        )

                        image_bytes = image_data[
                            "image"
                        ]

                        image_array = np.frombuffer(
                            image_bytes,
                            dtype=np.uint8
                        )

                        image = cv2.imdecode(
                            image_array,
                            cv2.IMREAD_COLOR
                        )

                        if image is not None:

                            self.process_image(
                                image,
                                pdf_path,
                                page_number
                            )

                    except Exception as error:
                        self.root.after(
                            0,
                            self.add_scan_error,
                            pdf_path,
                            "Не удалось извлечь изображение",
                            error
                        )

                # ------------------------------------------------
                # РЕНДЕР
                # ------------------------------------------------

                for scale in [
                    3.0,
                    4.0,
                    5.0
                ]:

                    if self.stop_requested:
                        break

                    matrix = fitz.Matrix(
                        scale,
                        scale
                    )

                    try:

                        pix = page.get_pixmap(
                            matrix=matrix,
                            alpha=False
                        )

                        image_array = np.frombuffer(
                            pix.samples,
                            dtype=np.uint8
                        )

                        image = image_array.reshape(
                            pix.height,
                            pix.width,
                            pix.n
                        )

                        if pix.n == 4:

                            image = cv2.cvtColor(
                                image,
                                cv2.COLOR_RGBA2BGR
                            )

                        else:

                            image = cv2.cvtColor(
                                image,
                                cv2.COLOR_RGB2BGR
                            )

                        self.process_image(
                            image,
                            pdf_path,
                            page_number
                        )

                    except Exception as error:
                        self.root.after(
                            0,
                            self.add_scan_error,
                            pdf_path,
                            f"Не удалось отрендерить страницу {page_number}",
                            error
                        )

                processed_pages += 1

                progress_value = (
                    processed_pages /
                    total_pages
                ) * 100

                self.root.after(
                    0,
                    self.progress.configure,
                    {"value": progress_value}
                )

            document.close()

        self.root.after(
            0,
            self.scan_finished
        )

    # ============================================================
    # ОБРАБОТКА ИЗОБРАЖЕНИЯ
    # ============================================================

    def process_image(
        self,
        image,
        pdf_path,
        page_number
    ):

        if self.stop_requested:
            return

        variants = [
            image
        ]

        gray = cv2.cvtColor(
            image,
            cv2.COLOR_BGR2GRAY
        )

        variants.append(
            gray
        )

        clahe = cv2.createCLAHE(
            clipLimit=2.0,
            tileGridSize=(8, 8)
        )

        enhanced = clahe.apply(
            gray
        )

        variants.append(
            enhanced
        )

        _, otsu = cv2.threshold(
            enhanced,
            0,
            255,
            cv2.THRESH_BINARY +
            cv2.THRESH_OTSU
        )

        variants.append(
            otsu
        )

        adaptive = cv2.adaptiveThreshold(
            enhanced,
            255,
            cv2.ADAPTIVE_THRESH_GAUSSIAN_C,
            cv2.THRESH_BINARY,
            31,
            11
        )

        variants.append(
            adaptive
        )

        for variant in variants:

            if self.stop_requested:
                return

            try:

                decoded = zxingcpp.read_barcodes(
                    variant
                )

                for barcode in decoded:

                    text = barcode.text

                    if not text:
                        continue

                    code = self.clean_code(
                        text
                    )

                    if not self.is_valid_marking_code(
                        code
                    ):
                        continue

                    self.root.after(
                        0,
                        self.add_result,
                        pdf_path,
                        page_number,
                        code
                    )

            except Exception as error:
                self.root.after(
                    0,
                    self.add_scan_error,
                    pdf_path,
                    "Ошибка распознавания штрихкода",
                    error
                )

    # ============================================================
    # ОЧИСТКА КОДА
    # ============================================================

    def clean_code(
        self,
        code
    ):

        if not code:
            return ""

        code = str(code).strip()

        # --------------------------------------------------------
        # НОРМАЛИЗАЦИЯ GS1
        # --------------------------------------------------------
        # В готовом коде Честного знака разделитель между AI
        # должен быть настоящим ASCII Group Separator (29).
        #
        # Если декодер вернул текстовую последовательность
        # "\x1d", заменяем её на настоящий символ chr(29).
        # --------------------------------------------------------

        code = code.replace(
            r"\x1d",
            chr(29)
        )

        code = code.replace(
            r"\u001d",
            chr(29)
        )

        # --------------------------------------------------------
        # УБИРАЕМ СКОБКИ AI
        # --------------------------------------------------------
        # (01)046... -> 01046...
        # (21)ABC... -> 21ABC...
        # (91)EE12   -> GS91EE12
        # (92)XYZ    -> GS92XYZ
        #
        # Важно: настоящий GS внутри уже отсканированного кода
        # не изменяется.
        # --------------------------------------------------------

        if code.startswith("(01)"):
            code = "01" + code[4:]

        code = code.replace(
            "(21)",
            "21",
            1
        )

        code = code.replace(
            "(91)",
            chr(29) + "91",
            1
        )

        code = code.replace(
            "(92)",
            chr(29) + "92",
            1
        )

        return code

    # ============================================================
    # ПРОВЕРКА
    # ============================================================

    def is_valid_marking_code(
        self,
        code
    ):

        if not code:
            return False

        normalized = code

        if not normalized.startswith("01"):
            return False

        gtin = normalized[2:16]

        return self.is_valid_gtin(gtin)

    def is_valid_gtin(self, gtin):

        if len(gtin) != 14 or not gtin.isdigit():
            return False

        check_sum = sum(
            int(digit) * (3 if index % 2 == 0 else 1)
            for index, digit in enumerate(reversed(gtin[:-1]))
        )
        check_digit = (10 - check_sum % 10) % 10

        return check_digit == int(gtin[-1])

    # ============================================================
    # ДОБАВЛЕНИЕ РЕЗУЛЬТАТА
    # ============================================================

    def add_result(
        self,
        pdf_path,
        page_number,
        code
    ):

        if code in self.seen_codes:
            return

        self.seen_codes.add(code)

        parsed = self.parse_marking_code(
            code
        )

        result = {
            "file": os.path.basename(pdf_path),
            "file_path": pdf_path,
            "page": page_number,
            "code": code,
            "gtin": parsed["gtin"],
            "serial": parsed["serial"],
            "crypto_91": parsed["crypto_91"],
            "crypto_92": parsed["crypto_92"],
            "crypto": parsed["crypto"]
        }

        self.results.append(
            result
        )

        self.file_stats[pdf_path] = (
            self.file_stats.get(
                pdf_path,
                0
            ) + 1
        )

        self.refresh_file_tree()
        self.update_code_count()

    # ============================================================
    # РАЗБОР КОДА
    # ============================================================

    def parse_marking_code(
        self,
        code
    ):

        result = {
            "gtin": "",
            "serial": "",
            "crypto_91": "",
            "crypto_92": "",
            "crypto": ""
        }

        if not code:
            return result

        raw = code

        if "(01)" in raw:

            raw = raw.replace(
                "(01)",
                "01"
            )

            raw = raw.replace(
                "(21)",
                "\x1d21"
            )

            raw = raw.replace(
                "(91)",
                "\x1d91"
            )

            raw = raw.replace(
                "(92)",
                "\x1d92"
            )

        # --------------------------------------------------------
        # GTIN
        # --------------------------------------------------------

        pos_01 = raw.find(
            "01"
        )

        if (
            pos_01 >= 0
            and len(raw) >= pos_01 + 16
        ):

            gtin = raw[
                pos_01 + 2:
                pos_01 + 16
            ]

            if gtin.isdigit():

                result["gtin"] = gtin

        # --------------------------------------------------------
        # AI 21
        # --------------------------------------------------------

        pos_21 = raw.find(
            "\x1d21"
        )

        if pos_21 >= 0:

            serial_start = (
                pos_21 + 3
            )

            serial_end = len(
                raw
            )

            for marker in [
                "\x1d91",
                "\x1d92",
                "\x1d93",
                "\x1d94",
                "\x1d95",
                "\x1d96",
                "\x1d97",
                "\x1d98",
                "\x1d99"
            ]:

                found = raw.find(
                    marker,
                    serial_start
                )

                if found >= 0:

                    serial_end = min(
                        serial_end,
                        found
                    )

            result["serial"] = raw[
                serial_start:
                serial_end
            ]

        else:

            if pos_01 >= 0:

                gtin_end = (
                    pos_01 + 16
                )

                pos_21_plain = raw.find(
                    "21",
                    gtin_end
                )

                if pos_21_plain >= 0:

                    serial_start = (
                        pos_21_plain + 2
                    )

                    serial_end = len(
                        raw
                    )

                    for marker in [
                        "91",
                        "92"
                    ]:

                        found = raw.find(
                            marker,
                            serial_start
                        )

                        if found >= 0:

                            serial_end = min(
                                serial_end,
                                found
                            )

                    result["serial"] = raw[
                        serial_start:
                        serial_end
                    ]

        # --------------------------------------------------------
        # AI 91
        # --------------------------------------------------------

        pos_91 = raw.find(
            "\x1d91"
        )

        if pos_91 >= 0:

            start = (
                pos_91 + 3
            )

            end = raw.find(
                "\x1d92",
                start
            )

            if end < 0:
                end = len(raw)

            result["crypto_91"] = raw[
                start:end
            ]

        else:

            pos_91_plain = raw.find(
                "91"
            )

            if pos_91_plain >= 0:

                start = (
                    pos_91_plain + 2
                )

                end = raw.find(
                    "92",
                    start
                )

                if end < 0:
                    end = len(raw)

                result["crypto_91"] = raw[
                    start:end
                ]

        # --------------------------------------------------------
        # AI 92
        # --------------------------------------------------------

        pos_92 = raw.find(
            "\x1d92"
        )

        if pos_92 >= 0:

            start = (
                pos_92 + 3
            )

            result["crypto_92"] = raw[
                start:
            ]

        else:

            pos_92_plain = raw.find(
                "92"
            )

            if pos_92_plain >= 0:

                start = (
                    pos_92_plain + 2
                )

                result["crypto_92"] = raw[
                    start:
                ]

        # --------------------------------------------------------
        # КРИПТОХВОСТ
        # --------------------------------------------------------

        crypto_parts = []

        if result["crypto_91"]:

            crypto_parts.append(
                "(91)" +
                result["crypto_91"]
            )

        if result["crypto_92"]:

            crypto_parts.append(
                "(92)" +
                result["crypto_92"]
            )

        result["crypto"] = "".join(
            crypto_parts
        )

        return result

    # ============================================================
    # ЛЕВОЕ ДЕРЕВО
    # ============================================================

    def refresh_file_tree(self):

        for item in self.file_tree.get_children():

            self.file_tree.delete(
                item
            )

        for pdf_path in self.pdf_files:

            name = os.path.basename(
                pdf_path
            )

            count = self.file_stats.get(pdf_path, 0)

            self.file_tree.insert(
                "",
                "end",
                iid=pdf_path,
                text=name,
                values=(count,)
            )

    # ============================================================
    # ВЫБОР PDF
    # ============================================================

    def on_file_select(
        self,
        event=None
    ):

        selected = (
            self.file_tree.selection()
        )

        if not selected:
            return

        pdf_path = selected[0]

        self.code_tree.delete(
            *self.code_tree.get_children()
        )

        counter = 1

        for item in self.results:

            if item["file_path"] != pdf_path:
                continue

            self.code_tree.insert(
                "",
                "end",
                iid=str(id(item)),
                values=(
                    counter,
                    item["page"],
                    item["code"]
                )
            )

            counter += 1

    # ============================================================
    # ВЫБОР КОДА
    # ============================================================

    def on_code_select(
        self,
        event=None
    ):

        selected = (
            self.code_tree.selection()
        )

        if selected:

            self.copy_button.config(
                state="normal"
            )

            self.details_button.config(
                state="normal"
            )

        else:

            self.copy_button.config(
                state="disabled"
            )

            self.details_button.config(
                state="disabled"
            )

    # ============================================================
    # ПОЛУЧЕНИЕ КОДА
    # ============================================================

    def get_selected_result(self):

        selected = (
            self.code_tree.selection()
        )

        if not selected:
            return None

        selected_id = selected[0]

        for item in self.results:

            if str(id(item)) == selected_id:

                return item

        return None

    # ============================================================
    # КОПИРОВАНИЕ
    # ============================================================

    def copy_selected_code(
        self,
        event=None
    ):

        item = self.get_selected_result()

        if not item:
            return

        self.root.clipboard_clear()

        self.root.clipboard_append(
            item["code"]
        )

        self.root.update()

        self.update_status(
            "Полный код скопирован"
        )

    # ============================================================
    # ПОДРОБНОСТИ
    # ============================================================

    def show_code_details(self):

        item = self.get_selected_result()

        if not item:
            return

        window = tk.Toplevel(
            self.root
        )

        window.title(
            "Подробности кода маркировки"
        )

        window.geometry(
            "900x650"
        )

        window.minsize(
            700,
            500
        )

        window.configure(
            bg=self.BG
        )

        window.transient(
            self.root
        )

        # --------------------------------------------------------
        # Заголовок
        # --------------------------------------------------------

        header = tk.Frame(
            window,
            bg=self.BG
        )

        header.pack(
            fill="x",
            padx=25,
            pady=(20, 12)
        )

        tk.Label(
            header,
            text="Подробности кода",
            bg=self.BG,
            fg=self.TEXT,
            font=("Segoe UI", 18, "bold")
        ).pack(
            anchor="w"
        )

        tk.Label(
            header,
            text="Информация, распознанная из Data Matrix",
            bg=self.BG,
            fg=self.TEXT_SECONDARY,
            font=("Segoe UI", 9)
        ).pack(
            anchor="w"
        )

        # --------------------------------------------------------
        # Контейнер
        # --------------------------------------------------------

        frame = tk.Frame(
            window,
            bg=self.PANEL,
            highlightbackground=self.BORDER,
            highlightthickness=1
        )

        frame.pack(
            fill="both",
            expand=True,
            padx=25,
            pady=(0, 20)
        )

        content = tk.Frame(
            frame,
            bg=self.PANEL
        )

        content.pack(
            fill="both",
            expand=True,
            padx=20,
            pady=20
        )

        # --------------------------------------------------------
        # Поля
        # --------------------------------------------------------

        self.create_detail_field(
            content,
            "GTIN",
            item["gtin"]
        )

        self.create_detail_field(
            content,
            "Серийный номер (AI 21)",
            item["serial"]
        )

        self.create_detail_field(
            content,
            "Криптохвост AI 91",
            item["crypto_91"]
        )

        self.create_detail_field(
            content,
            "Криптохвост AI 92",
            item["crypto_92"]
        )

        # --------------------------------------------------------
        # Полный код
        # --------------------------------------------------------

        tk.Label(
            content,
            text="Полный код маркировки",
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 10, "bold")
        ).pack(
            anchor="w",
            pady=(8, 5)
        )

        text_frame = tk.Frame(
            content,
            bg=self.PANEL_2,
            highlightbackground=self.BORDER,
            highlightthickness=1
        )

        text_frame.pack(
            fill="both",
            expand=True
        )

        full_code_text = tk.Text(
            text_frame,
            wrap="none",
            bg=self.PANEL_2,
            fg=self.TEXT,
            insertbackground=self.TEXT,
            selectbackground=self.SELECTED,
            selectforeground="white",
            relief="flat",
            borderwidth=0,
            font=("Consolas", 10),
            padx=10,
            pady=10
        )

        full_code_text.pack(
            side="left",
            fill="both",
            expand=True
        )

        scroll = ttk.Scrollbar(
            text_frame,
            orient="vertical",
            command=full_code_text.yview,
            style="Dark.Vertical.TScrollbar"
        )

        scroll.pack(
            side="right",
            fill="y"
        )

        full_code_text.configure(
            yscrollcommand=scroll.set
        )

        full_code_text.insert(
            "1.0",
            item["code"]
        )

        full_code_text.config(
            state="disabled"
        )

        # --------------------------------------------------------
        # Кнопки
        # --------------------------------------------------------

        buttons = tk.Frame(
            content,
            bg=self.PANEL
        )

        buttons.pack(
            fill="x",
            pady=(15, 0)
        )

        def copy_value(
            value,
            message
        ):

            self.root.clipboard_clear()

            self.root.clipboard_append(
                value
            )

            self.root.update()

            self.update_status(
                message
            )

        ttk.Button(
            buttons,
            text="Копировать GTIN",
            command=lambda: copy_value(
                item["gtin"],
                "GTIN скопирован"
            ),
            style="Dark.TButton"
        ).pack(
            side="left",
            padx=(0, 5)
        )

        ttk.Button(
            buttons,
            text="Копировать серийный №",
            command=lambda: copy_value(
                item["serial"],
                "Серийный номер скопирован"
            ),
            style="Dark.TButton"
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            buttons,
            text="Копировать криптохвост",
            command=lambda: copy_value(
                item["crypto"],
                "Криптохвост скопирован"
            ),
            style="Dark.TButton"
        ).pack(
            side="left",
            padx=5
        )

        ttk.Button(
            buttons,
            text="Закрыть",
            command=window.destroy,
            style="Accent.TButton"
        ).pack(
            side="right"
        )

    # ============================================================
    # ПОЛЕ ПОДРОБНОСТЕЙ
    # ============================================================

    def create_detail_field(
        self,
        parent,
        title,
        value
    ):

        tk.Label(
            parent,
            text=title,
            bg=self.PANEL,
            fg=self.TEXT,
            font=("Segoe UI", 10, "bold")
        ).pack(
            anchor="w",
            pady=(0, 4)
        )

        entry = tk.Entry(
            parent,
            bg=self.PANEL_2,
            fg=self.TEXT,
            insertbackground=self.TEXT,
            readonlybackground=self.PANEL_2,
            relief="flat",
            font=("Consolas", 10),
            highlightbackground=self.BORDER,
            highlightthickness=1
        )

        entry.pack(
            fill="x",
            pady=(0, 10),
            ipady=6
        )

        entry.insert(
            0,
            value
        )

        entry.config(
            state="readonly"
        )

    # ============================================================
    # СТАТИСТИКА
    # ============================================================

    def update_code_count(self):

        count = len(
            self.results
        )

        self.code_count_value.config(
            text=str(count)
        )

    # ============================================================
    # СТАТУС
    # ============================================================

    def update_status(
        self,
        text
    ):

        self.status_label.config(
            text=text
        )

        self.status_value.config(
            text=text[:22]
        )

    # ============================================================
    # ОШИБКА
    # ============================================================

    def show_error(
        self,
        title,
        message
    ):

        messagebox.showerror(
            title,
            message
        )

    # ============================================================
    # ЗАВЕРШЕНИЕ
    # ============================================================

    def scan_finished(self):

        self.scanning = False

        self.select_button.config(
            state="normal"
        )

        self.stop_button.config(
            state="disabled"
        )

        if self.results:

            self.excel_button.config(
                state="normal"
            )

            self.csv_button.config(
                state="normal"
            )

            self.print_button.config(
                state="normal"
            )

        self.progress["value"] = 100

        if self.stop_requested:

            self.update_status(
                "Сканирование остановлено"
            )

            self.header_status.config(
                text="● Остановлено",
                fg=self.DANGER
            )

            self.status_value.config(
                text="Остановлено"
            )

        else:

            error_suffix = (
                f" Ошибок: {len(self.scan_errors)}"
                if self.scan_errors else ""
            )

            self.update_status(
                f"Готово. Найдено кодов: {len(self.results)}.{error_suffix}"
            )

            self.header_status.config(
                text="● Готово",
                fg=self.SUCCESS
            )

            self.status_value.config(
                text="Готово"
            )

        self.refresh_file_tree()

    def add_scan_error(self, pdf_path, context, error):

        message = f"{os.path.basename(pdf_path)}: {context}: {error}"

        if message not in self.scan_errors:
            self.scan_errors.append(message)

    # ============================================================
    # ОЧИСТКА
    # ============================================================

    def clear_results(
        self,
        keep_files=False
    ):

        self.results.clear()

        self.file_stats.clear()

        self.seen_codes.clear()

        self.scan_errors.clear()

        self.file_tree.delete(
            *self.file_tree.get_children()
        )

        self.code_tree.delete(
            *self.code_tree.get_children()
        )

        self.code_count_value.config(
            text="0"
        )

        self.progress["value"] = 0

        self.copy_button.config(
            state="disabled"
        )

        self.details_button.config(
            state="disabled"
        )

        self.excel_button.config(
            state="disabled"
        )

        self.csv_button.config(
            state="disabled"
        )

        self.print_button.config(
            state="disabled"
        )

        if not keep_files:

            self.pdf_files = []

            self.pdf_count_value.config(
                text="0"
            )

        self.update_status(
            "Готов"
        )

        self.header_status.config(
            text="● Готов",
            fg=self.SUCCESS
        )

        self.status_value.config(
            text="Готов"
        )

    # ============================================================
    # ПЕЧАТЬ ЭТИКЕТОК ZPL
    # ============================================================

    def get_selected_results(self):

        selected_ids = set(self.code_tree.selection())

        return [
            item
            for item in self.results
            if str(id(item)) in selected_ids
        ]

    def zpl_escape(self, value):

        escaped = []

        for byte in value.encode("utf-8"):

            char = chr(byte)

            if 32 <= byte <= 126 and char not in "^~\\":
                escaped.append(char)
            else:
                escaped.append(f"\\{byte:02X}")

        return "".join(escaped)

    def build_zpl_label(self, item, width_mm, height_mm, dots_per_mm):

        width = width_mm * dots_per_mm
        height = height_mm * dots_per_mm
        margin = 3 * dots_per_mm
        matrix_size = max(3, min(8, int(height / 65)))
        text_y = margin + matrix_size * 30 + (2 * dots_per_mm)

        code = self.zpl_escape(item["code"])
        gtin = self.zpl_escape(item["gtin"] or "—")
        serial = self.zpl_escape(item["serial"] or "—")

        return (
            "^XA\n"
            f"^PW{width}\n"
            f"^LL{height}\n"
            "^CI28\n"
            f"^FO{margin},{margin}^BXN,{matrix_size},200^FH\\^FD{code}^FS\n"
            f"^FO{margin},{text_y}^A0N,{2 * dots_per_mm},{2 * dots_per_mm}^FDGTIN: {gtin}^FS\n"
            f"^FO{margin},{text_y + 3 * dots_per_mm}^A0N,{2 * dots_per_mm},{2 * dots_per_mm}^FDSN: {serial}^FS\n"
            "^XZ\n"
        )

    def print_zpl(self):

        selected = self.get_selected_results()

        if selected:
            use_selected = messagebox.askyesno(
                "Печать этикеток",
                f"Напечатать выбранные коды ({len(selected)})?\n"
                "Нажмите «Нет», чтобы напечатать все найденные коды."
            )
            items = selected if use_selected else self.results
        else:
            items = self.results

        if not items:
            return

        width_mm = simpledialog.askinteger(
            "Размер этикетки",
            "Ширина этикетки, мм:",
            parent=self.root,
            initialvalue=58,
            minvalue=20,
            maxvalue=200
        )

        if width_mm is None:
            return

        height_mm = simpledialog.askinteger(
            "Размер этикетки",
            "Высота этикетки, мм:",
            parent=self.root,
            initialvalue=40,
            minvalue=20,
            maxvalue=200
        )

        if height_mm is None:
            return

        zpl = "".join(
            self.build_zpl_label(item, width_mm, height_mm, 8)
            for item in items
        )

        try:
            import win32print

            printer_name = win32print.GetDefaultPrinter()

            if not messagebox.askyesno(
                "Подтвердите печать",
                f"Отправить {len(items)} этикеток на принтер:\n{printer_name}?\n\n"
                "Принтер должен поддерживать ZPL (например, Zebra)."
            ):
                return

            printer = win32print.OpenPrinter(printer_name)

            try:
                job = win32print.StartDocPrinter(
                    printer,
                    1,
                    ("Chestny Znak labels", None, "RAW")
                )
                win32print.StartPagePrinter(printer)
                win32print.WritePrinter(printer, zpl.encode("utf-8"))
                win32print.EndPagePrinter(printer)
                win32print.EndDocPrinter(printer)
            finally:
                win32print.ClosePrinter(printer)

            self.update_status(f"Отправлено на печать: {len(items)}")

        except ImportError:
            self.save_zpl_file(zpl)

        except Exception as error:
            messagebox.showerror("Ошибка печати", str(error))

    def save_zpl_file(self, zpl):

        filename = filedialog.asksaveasfilename(
            title="Сохранить ZPL для термопринтера",
            defaultextension=".zpl",
            filetypes=[
                ("ZPL-файл", "*.zpl"),
                ("Все файлы", "*.*")
            ]
        )

        if not filename:
            return

        try:
            with open(filename, "w", encoding="utf-8", newline="") as file:
                file.write(zpl)

            messagebox.showinfo(
                "ZPL сохранён",
                "Модуль pywin32 не установлен, поэтому этикетки сохранены "
                "в ZPL-файл. Его можно отправить на принтер Zebra."
            )

        except Exception as error:
            messagebox.showerror("Ошибка сохранения ZPL", str(error))

    # ============================================================
    # EXCEL
    # ============================================================

    def save_excel(self):

        if not self.results:
            return

        filename = filedialog.asksaveasfilename(
            title="Сохранить результаты в Excel",
            defaultextension=".xlsx",
            filetypes=[
                ("Excel-файл", "*.xlsx"),
                ("Все файлы", "*.*")
            ]
        )

        if not filename:
            return

        try:

            workbook = Workbook()

            sheet = workbook.active

            sheet.title = "Коды"

            headers = [
                "№",
                "PDF-файл",
                "Страница",
                "GTIN",
                "Серийный номер",
                "Криптохвост AI 91",
                "Криптохвост AI 92",
                "Криптохвост",
                "Полный код"
            ]

            sheet.append(
                headers
            )

            for index, item in enumerate(
                self.results,
                start=1
            ):

                sheet.append([
                    index,
                    item["file"],
                    item["page"],
                    item["gtin"],
                    item["serial"],
                    item["crypto_91"],
                    item["crypto_92"],
                    item["crypto"],
                    item["code"]
                ])

            sheet.auto_filter.ref = (
                f"A1:I{sheet.max_row}"
            )

            sheet.freeze_panes = "A2"

            widths = {
                "A": 8,
                "B": 45,
                "C": 12,
                "D": 18,
                "E": 30,
                "F": 35,
                "G": 35,
                "H": 70,
                "I": 100
            }

            for column, width in widths.items():

                sheet.column_dimensions[
                    column
                ].width = width

            workbook.save(
                filename
            )

            messagebox.showinfo(
                "Готово",
                f"Excel сохранён:\n{filename}"
            )

        except Exception as e:

            messagebox.showerror(
                "Ошибка сохранения Excel",
                str(e)
            )

    # ============================================================
    # CSV
    # ============================================================

    def save_csv(self):

        if not self.results:
            return

        filename = filedialog.asksaveasfilename(
            title="Сохранить результаты в CSV",
            defaultextension=".csv",
            filetypes=[
                ("CSV-файл", "*.csv"),
                ("Все файлы", "*.*")
            ]
        )

        if not filename:
            return

        try:

            with open(
                filename,
                "w",
                newline="",
                encoding="utf-8-sig"
            ) as file:

                writer = csv.writer(
                    file,
                    delimiter=";",
                    quoting=csv.QUOTE_MINIMAL
                )

                writer.writerow([
                    "№",
                    "PDF-файл",
                    "Страница",
                    "GTIN",
                    "Серийный номер",
                    "Криптохвост AI 91",
                    "Криптохвост AI 92",
                    "Криптохвост",
                    "Полный код"
                ])

                for index, item in enumerate(
                    self.results,
                    start=1
                ):

                    writer.writerow([
                        index,
                        item["file"],
                        item["page"],
                        item["gtin"],
                        item["serial"],
                        item["crypto_91"],
                        item["crypto_92"],
                        item["crypto"],
                        item["code"]
                    ])

            messagebox.showinfo(
                "Готово",
                f"CSV сохранён:\n{filename}"
            )

        except Exception as e:

            messagebox.showerror(
                "Ошибка сохранения CSV",
                str(e)
            )


# ================================================================
# ЗАПУСК
# ================================================================

if __name__ == "__main__":

    root = tk.Tk()

    app = ChestnyZnakScanner(
        root
    )

    root.mainloop()
